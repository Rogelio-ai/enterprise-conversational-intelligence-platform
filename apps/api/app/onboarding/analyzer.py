"""Non-mutating analysis and preview for the Stage 0 onboarding workbook."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Any
import unicodedata
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from app.onboarding.contract import (
    CONTRACT_VERSION,
    MAX_CAPTURE_ROWS,
    ONBOARDING_CONTRACT,
    Field,
    Group,
)
from app.onboarding.official_contract import (
    CAPTURE_STATUSES,
    OFFICIAL_BY_NAME,
    OFFICIAL_METADATA_KEYS,
    OFFICIAL_MAX_CAPTURE_ROWS,
    OFFICIAL_PROFILE,
    OFFICIAL_REQUIRED_METADATA,
    OFFICIAL_SHEETS,
    OFFICIAL_WORKBOOK_SHEETS,
    OfficialSheet,
)


MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ARCHIVE_MEMBERS = 2_048
TEMPLATE_KIND = "stage0-real-restaurant-capture"
_SYSTEM_SHEETS = ("Instructions", "Metadata", "_Lists")
_CODE_PATTERN = re.compile(r"^[A-Z0-9][A-Z0-9_-]{0,63}$")
_SLUG_PATTERN = re.compile(r"^[a-z0-9][a-z0-9-]{0,99}$")
_EMAIL_PATTERN = re.compile(r"^[^\s@]+@[^\s@]+$")
_LANGUAGE_PATTERN = re.compile(r"^[A-Za-z]{2,8}(?:-[A-Za-z0-9]{1,8})*$")
_UNIT_PATTERN = re.compile(r"^[A-Z][A-Z0-9_]{0,31}$")
_DECIMAL_SHAPE = re.compile(r"decimal\((\d+),(\d+)\)")
_LEGACY_HEADERS = {
    "categories": (
        "organization_code", "category_name", "parent_category_name",
        "display_order", "status",
    ),
    "products": (
        "product_key", "organization_code", "category_name",
        "name", "description", "status",
    ),
}
_LEGACY_CATEGORY_FIELDS = (
    Field("organization_code", "Código organización", "text", True, "", reference="restaurant_profile.organization_code"),
    Field("category_name", "Categoría", "text", True, ""),
    Field("parent_category_name", "Categoría padre", "text", False, "", reference="categories.category_name"),
    Field("display_order", "Orden", "integer", False, "", format=">=0"),
    Field("status", "Estado", "enum", True, "", allowed_values=("ACTIVE", "INACTIVE")),
)
_LEGACY_PRODUCT_FIELDS = (
    Field("product_key", "Clave producto", "text", True, "", format="^[A-Z][A-Z0-9_-]{0,63}$"),
    Field("organization_code", "Código organización", "text", True, "", reference="restaurant_profile.organization_code"),
    Field("category_name", "Categoría", "text", False, "", reference="categories.category_name"),
    Field("name", "Producto", "text", True, ""),
    Field("description", "Descripción", "text", False, ""),
    Field("status", "Estado", "enum", True, "", allowed_values=("ACTIVE", "INACTIVE")),
)


@dataclass(frozen=True, slots=True)
class AnalysisIssue:
    code: str
    message: str
    group: str | None = None
    sheet: str | None = None
    row: int | None = None
    field: str | None = None
    value: str | None = None
    severity: str = "ERROR"


@dataclass(frozen=True, slots=True)
class AnalyzedRow:
    group: str
    sheet: str
    row: int
    business_key: dict[str, Any]
    values: dict[str, Any]
    errors: tuple[AnalysisIssue, ...]

    @property
    def valid(self) -> bool:
        return not self.errors


@dataclass(frozen=True, slots=True)
class GroupSummary:
    group: str
    sheet: str
    populated_rows: int
    valid_rows: int
    invalid_rows: int


@dataclass(frozen=True, slots=True)
class AnalysisResult:
    contract_version: str | None
    status: str
    total_populated_rows: int
    valid_rows: int
    invalid_rows: int
    warning_count: int
    persisted_rows: int
    groups: tuple[GroupSummary, ...]
    rows: tuple[AnalyzedRow, ...]
    errors: tuple[AnalysisIssue, ...]
    warnings: tuple[AnalysisIssue, ...]
    profile: str = "legacy-generated"
    metadata: dict[str, Any] | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return a deterministic, JSON-compatible preview representation."""

        return asdict(self)


@dataclass(slots=True)
class _WorkingRow:
    group: Group
    row: int
    values: dict[str, Any]
    errors: list[AnalysisIssue]


def analyze_xlsx(content: bytes, *, filename: str | None = None) -> AnalysisResult:
    """Analyze an official workbook without accepting any persistence dependency."""

    boundary_errors: list[AnalysisIssue] = []
    if filename is not None and not filename.strip().lower().endswith(".xlsx"):
        boundary_errors.append(_workbook_issue("INVALID_FILE_TYPE", "Only .xlsx workbooks are supported."))
    if not content:
        boundary_errors.append(_workbook_issue("EMPTY_FILE", "The workbook file is empty."))
    elif len(content) > MAX_XLSX_BYTES:
        boundary_errors.append(
            _workbook_issue("FILE_TOO_LARGE", f"The workbook exceeds the {MAX_XLSX_BYTES}-byte limit.")
        )
    elif not is_zipfile(BytesIO(content)):
        boundary_errors.append(_workbook_issue("INVALID_XLSX", "The file is not a valid XLSX/OOXML archive."))
    else:
        archive_error = _validate_archive_boundary(content)
        if archive_error is not None:
            boundary_errors.append(archive_error)
    if boundary_errors:
        return _empty_result(None, boundary_errors)

    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_vba=False)
    except Exception as exc:
        return _empty_result(
            None,
            [_workbook_issue("INVALID_XLSX", f"The XLSX workbook is malformed ({type(exc).__name__}).")],
        )

    try:
        if _is_official_profile(workbook):
            return _analyze_official(workbook)
        version, structural_errors, warnings = _validate_structure(workbook)
        if structural_errors:
            return _empty_result(version, structural_errors, warnings)

        working_rows = _parse_rows(workbook)
        _validate_duplicates(working_rows)
        row_contexts = _validate_references(working_rows)
        _validate_dataset_rules(working_rows, row_contexts)
        return _build_result(version, working_rows, warnings)
    finally:
        workbook.close()


def _is_official_profile(workbook: Any) -> bool:
    if "Metadata" not in workbook.sheetnames:
        return any(name in workbook.sheetnames for name in OFFICIAL_BY_NAME)
    return _trimmed_headers(workbook["Metadata"]) == (
        "field", "value", "required", "notes",
    ) or any(name in workbook.sheetnames for name in OFFICIAL_BY_NAME)


def _analyze_official(workbook: Any) -> AnalysisResult:
    structural_errors, warnings = _validate_official_structure(workbook)
    metadata, metadata_errors = _parse_official_metadata(workbook)
    version = metadata.get("contract_version")
    if structural_errors:
        return _empty_result(
            str(version) if version is not None else None,
            [*structural_errors, *metadata_errors],
            warnings,
            profile=OFFICIAL_PROFILE,
            metadata=metadata,
            groups=tuple(
                GroupSummary(sheet.key, sheet.name, 0, 0, 0)
                for sheet in OFFICIAL_SHEETS
            ),
        )
    rows = _parse_official_rows(workbook)
    _validate_official_duplicates(rows)
    _validate_official_references(rows)
    _validate_official_dataset(rows, metadata)
    return _build_official_result(
        str(version), rows, warnings, metadata, global_errors=metadata_errors,
    )


def _validate_official_structure(
    workbook: Any,
) -> tuple[list[AnalysisIssue], list[AnalysisIssue]]:
    errors: list[AnalysisIssue] = []
    warnings: list[AnalysisIssue] = []
    expected = tuple(OFFICIAL_WORKBOOK_SHEETS)
    for name in expected:
        if name not in workbook.sheetnames:
            errors.append(_workbook_issue(
                "MISSING_SHEET", f'Required official sheet "{name}" is missing.', sheet=name,
            ))
    for name in workbook.sheetnames:
        if name not in expected:
            errors.append(_workbook_issue(
                "UNEXPECTED_SHEET", f'Sheet "{name}" is not part of the official v1 profile.', sheet=name,
            ))
    if tuple(workbook.sheetnames) != expected and not any(
        issue.code in {"MISSING_SHEET", "UNEXPECTED_SHEET"} for issue in errors
    ):
        errors.append(_workbook_issue(
            "INVALID_SHEET_ORDER", "Official workbook sheets must retain their approved order.",
        ))
    if "Instructions" in workbook.sheetnames:
        if workbook["Instructions"]["B3"].value != CONTRACT_VERSION:
            errors.append(_workbook_issue(
                "INVALID_TEMPLATE_IDENTITY",
                f'Instructions contract must be "{CONTRACT_VERSION}".',
                sheet="Instructions", row=3, field="contract_version",
            ))
    if "Metadata" in workbook.sheetnames and _trimmed_headers(workbook["Metadata"]) != (
        "field", "value", "required", "notes",
    ):
        errors.append(_workbook_issue(
            "INVALID_METADATA_HEADERS",
            'Metadata headers must be exactly "field", "value", "required", and "notes".',
            sheet="Metadata", row=1,
        ))
    for spec in OFFICIAL_SHEETS:
        if spec.name not in workbook.sheetnames:
            continue
        sheet = workbook[spec.name]
        headers = _trimmed_headers(sheet)
        if headers != spec.headers:
            duplicates = sorted({h for h in headers if h is not None and headers.count(h) > 1})
            for header in duplicates:
                errors.append(_workbook_issue(
                    "DUPLICATE_HEADER", f'Duplicate header "{header}".',
                    group=spec.key, sheet=spec.name, row=1, field=str(header),
                ))
            missing = [header for header in spec.headers if header not in headers]
            for header in missing:
                errors.append(_workbook_issue(
                    "MISSING_HEADER", f'Required official column "{header}" is missing.',
                    group=spec.key, sheet=spec.name, row=1, field=header,
                ))
            if not missing and not duplicates:
                errors.append(_workbook_issue(
                    "INVALID_HEADER_CONTRACT",
                    "Columns must match the official contract exactly and in order.",
                    group=spec.key, sheet=spec.name, row=1,
                ))
        if sheet.max_column > len(spec.headers):
            for row_number in _official_populated_row_numbers(sheet):
                if any(
                    sheet.cell(row_number, column).data_type == "f"
                    or _blank_to_none(sheet.cell(row_number, column).value) is not None
                    for column in range(len(spec.headers) + 1, sheet.max_column + 1)
                ):
                    errors.append(_workbook_issue(
                        "UNSUPPORTED_COLUMN_DATA", "Data exists outside official columns.",
                        group=spec.key, sheet=spec.name, row=row_number,
                    ))
    return errors, warnings


def _parse_official_metadata(
    workbook: Any,
) -> tuple[dict[str, Any], list[AnalysisIssue]]:
    metadata: dict[str, Any] = {}
    errors: list[AnalysisIssue] = []
    if "Metadata" not in workbook.sheetnames:
        return metadata, errors
    sheet = workbook["Metadata"]
    seen: set[str] = set()
    flags: dict[str, str] = {}
    for row_number in _populated_row_numbers(sheet):
        field = _normalized_text(sheet.cell(row_number, 1).value)
        value = _normalized_value(sheet.cell(row_number, 2).value)
        required = _normalized_text(sheet.cell(row_number, 3).value)
        if field is None:
            errors.append(_workbook_issue(
                "INVALID_METADATA", "Metadata field must be non-empty text.",
                sheet="Metadata", row=row_number, field="field",
            ))
            continue
        if field in seen:
            errors.append(_workbook_issue(
                "DUPLICATE_METADATA_KEY", f'Duplicate Metadata field "{field}".',
                sheet="Metadata", row=row_number, field=field,
            ))
            continue
        seen.add(field)
        if field not in OFFICIAL_METADATA_KEYS:
            errors.append(_workbook_issue(
                "UNKNOWN_METADATA_KEY", f'Metadata field "{field}" is not defined by v1.',
                sheet="Metadata", row=row_number, field=field,
            ))
        if required not in {"YES", "NO"}:
            errors.append(_workbook_issue(
                "INVALID_METADATA_REQUIRED", 'Metadata required must be "YES" or "NO".',
                sheet="Metadata", row=row_number, field=field, value=required,
            ))
        else:
            flags[field] = required
        metadata[field] = value
    for field in OFFICIAL_METADATA_KEYS:
        if field not in metadata:
            errors.append(_workbook_issue(
                "MISSING_METADATA_KEY", f'Metadata field "{field}" is required by the profile.',
                sheet="Metadata", field=field,
            ))
            continue
        expected = "YES" if field in OFFICIAL_REQUIRED_METADATA else "NO"
        if flags.get(field) != expected:
            errors.append(_workbook_issue(
                "METADATA_REQUIRED_CONFLICT",
                f'Metadata field "{field}" must declare required={expected}.',
                sheet="Metadata", field=field, value=flags.get(field),
            ))
        if field in OFFICIAL_REQUIRED_METADATA and metadata[field] is None:
            errors.append(_workbook_issue(
                "REQUIRED_METADATA_VALUE", f'Metadata value "{field}" is required.',
                sheet="Metadata", field=field,
            ))
    if metadata.get("contract_version") != CONTRACT_VERSION:
        errors.append(_workbook_issue(
            "UNSUPPORTED_CONTRACT_VERSION",
            f'Metadata.contract_version must be "{CONTRACT_VERSION}".',
            sheet="Metadata", field="contract_version", value=metadata.get("contract_version"),
        ))
    if metadata.get("dataset_kind") != "REAL_PREPILOT":
        errors.append(_workbook_issue(
            "INVALID_DATASET_KIND", 'Metadata.dataset_kind must be "REAL_PREPILOT".',
            sheet="Metadata", field="dataset_kind", value=metadata.get("dataset_kind"),
        ))
    capture_status = metadata.get("capture_status")
    if capture_status not in CAPTURE_STATUSES:
        errors.append(_workbook_issue(
            "INVALID_METADATA_STATUS", "Metadata.capture_status is not an approved capture state.",
            sheet="Metadata", field="capture_status", value=capture_status,
        ))
    confirmed_at = metadata.get("confirmed_at")
    if confirmed_at is not None:
        try:
            _parse_official_instant(confirmed_at)
        except ValueError:
            errors.append(_workbook_issue(
                "INVALID_METADATA_VALUE",
                "Metadata.confirmed_at must be an ISO 8601 value with timezone.",
                sheet="Metadata", field="confirmed_at", value=confirmed_at,
            ))
    return metadata, errors


def _normalized_text(value: Any) -> str | None:
    if not isinstance(value, str):
        return None if value is None else str(value)
    normalized = unicodedata.normalize("NFC", value.strip())
    return normalized or None


def _normalized_value(value: Any) -> Any:
    return _normalized_text(value) if isinstance(value, str) else value


def _parse_official_rows(workbook: Any) -> list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]]:
    rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]] = []
    for spec in OFFICIAL_SHEETS:
        sheet = workbook[spec.name]
        allowed = dict(spec.allowed)
        for row_number in _official_populated_row_numbers(sheet):
            values: dict[str, Any] = {}
            issues: list[AnalysisIssue] = []
            for column, field in enumerate(spec.headers, start=1):
                cell = sheet.cell(row_number, column)
                if cell.data_type == "f":
                    values[field] = None
                    issues.append(_official_row_issue(
                        spec, row_number, "FORMULA_NOT_ALLOWED",
                        "Formulas are not authoritative onboarding data.", field=field,
                    ))
                else:
                    values[field] = _normalized_value(_blank_to_none(cell.value))
            if row_number > OFFICIAL_MAX_CAPTURE_ROWS + 1:
                issues.append(_official_row_issue(
                    spec, row_number, "ROW_LIMIT_EXCEEDED",
                    f"Rows beyond {OFFICIAL_MAX_CAPTURE_ROWS + 1} are not supported.",
                ))
            if values.get("capture_status") != "NO APLICA":
                for field in spec.required:
                    if values.get(field) is None:
                        issues.append(_official_row_issue(
                            spec, row_number, "REQUIRED_FIELD", "A required value is missing.", field=field,
                        ))
            for field, choices in allowed.items():
                value = values.get(field)
                if value is not None and value not in choices:
                    issues.append(_official_row_issue(
                        spec, row_number, "INVALID_CONTROLLED_VALUE",
                        f'Value must be one of: {", ".join(choices)}.', field=field, value=value,
                    ))
            for field in spec.integer_fields:
                value = values.get(field)
                if value is not None:
                    try:
                        values[field] = _parse_integer(value)
                    except ValueError:
                        issues.append(_official_row_issue(
                            spec, row_number, "INVALID_FIELD_VALUE", "Value must be an integer.",
                            field=field, value=value,
                        ))
            for field in spec.decimal_fields:
                value = values.get(field)
                if value is not None:
                    try:
                        parsed = _parse_decimal(value)
                        if Decimal(parsed) < 0:
                            raise ValueError
                        values[field] = parsed
                    except (InvalidOperation, TypeError, ValueError):
                        issues.append(_official_row_issue(
                            spec, row_number, "INVALID_FIELD_VALUE",
                            "Value must be a non-negative finite decimal.", field=field, value=value,
                        ))
            for field in ("valid_from", "valid_to"):
                value = values.get(field)
                if value is not None:
                    try:
                        values[field] = _parse_official_instant(value)
                    except ValueError:
                        issues.append(_official_row_issue(
                            spec, row_number, "INVALID_FIELD_VALUE",
                            "Value must be ISO 8601 with timezone.", field=field, value=value,
                        ))
            for field in ("email", "email_optional"):
                value = values.get(field)
                if value is not None and not _EMAIL_PATTERN.fullmatch(str(value)):
                    issues.append(_official_row_issue(
                        spec, row_number, "INVALID_FIELD_VALUE", "Value must be an email address.",
                        field=field, value=value,
                    ))
            if spec.name == "UOM_Conversions":
                if values.get("from_uom") == values.get("to_uom"):
                    issues.append(_official_row_issue(
                        spec, row_number, "INVALID_FIELD_VALUE",
                        "from_uom and to_uom must differ.", field="to_uom", value=values.get("to_uom"),
                    ))
                factor = values.get("factor")
                if factor is not None:
                    try:
                        invalid_factor = Decimal(str(factor)) <= 0
                    except (InvalidOperation, ValueError):
                        invalid_factor = False  # The generic decimal validation already reported it.
                    if invalid_factor:
                        issues.append(_official_row_issue(
                            spec, row_number, "INVALID_FIELD_VALUE",
                            "factor must be positive.", field="factor", value=factor,
                        ))
            if spec.name == "Modifier_Groups":
                minimum, maximum = values.get("min_select"), values.get("max_select")
                if minimum is not None and maximum is not None and minimum > maximum:
                    issues.append(_official_row_issue(
                        spec, row_number, "INVALID_FIELD_VALUE",
                        "min_select cannot exceed max_select.", field="max_select", value=maximum,
                    ))
            rows.append((spec, row_number, values, issues))
    return rows


def _parse_official_instant(value: Any) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        text = value.strip()
        parsed = datetime.fromisoformat(text[:-1] + "+00:00" if text.endswith("Z") else text)
    else:
        raise ValueError
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError
    return parsed.isoformat()


def _official_populated_row_numbers(sheet: Any) -> tuple[int, ...]:
    """Return the approved 60-row capture area, excluding immutable footer guidance."""

    return tuple(
        sorted({
            cell.row
            for cell in sheet._cells.values()
            if 2 <= cell.row <= OFFICIAL_MAX_CAPTURE_ROWS + 1
            and (cell.data_type == "f" or _blank_to_none(cell.value) is not None)
        })
    )


def _validate_official_duplicates(
    rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]],
) -> None:
    seen: dict[tuple[str, tuple[Any, ...]], list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]]] = {}
    for row in rows:
        spec, _, values, _ = row
        if values.get("capture_status") == "NO APLICA":
            continue
        key = tuple(values.get(field) for field in spec.business_key)
        if all(value is not None for value in key):
            seen.setdefault((spec.key, key), []).append(row)
    for duplicates in seen.values():
        if len(duplicates) > 1:
            for spec, row_number, _, issues in duplicates:
                issues.append(_official_row_issue(
                    spec, row_number, "DUPLICATE_BUSINESS_KEY",
                    "The normalized business key is duplicated in this sheet.",
                ))


def _validate_official_references(
    rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]],
) -> None:
    indexes: dict[
        tuple[str, str],
        dict[Any, list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]]],
    ] = {}
    for row in rows:
        spec, _, values, _ = row
        if values.get("capture_status") == "NO APLICA":
            continue
        for field in spec.headers:
            value = values.get(field)
            if value is not None:
                index = indexes.setdefault((spec.name, field), {})
                index.setdefault(value, []).append(row)
    derived_locations = _official_row_locations(rows, indexes)
    for spec, row_number, values, issues in rows:
        if values.get("capture_status") == "NO APLICA":
            continue
        for field, target_sheet, target_field in spec.references:
            value = values.get(field)
            if value is None:
                continue
            candidates = list(indexes.get((target_sheet, target_field), {}).get(value, ()))
            source_location = derived_locations.get((spec.name, row_number))
            if source_location is not None:
                scoped = [
                    candidate for candidate in candidates
                    if derived_locations.get((candidate[0].name, candidate[1])) in {
                        None, source_location,
                    }
                ]
                candidates = scoped
            count = len(candidates)
            if count != 1:
                issues.append(_official_row_issue(
                    spec, row_number,
                    "UNRESOLVED_REFERENCE" if count == 0 else "AMBIGUOUS_REFERENCE",
                    f'{field} must resolve exactly once in {target_sheet}.{target_field}.',
                    field=field, value=value,
                ))
    category_parents = {
        values.get("category_code"): values.get("parent_category_code")
        for spec, _, values, _ in rows
        if spec.name == "Categories" and values.get("capture_status") != "NO APLICA"
    }
    for spec, row_number, values, issues in rows:
        if spec.name != "Categories" or values.get("category_code") is None:
            continue
        origin = values["category_code"]
        cursor = category_parents.get(origin)
        visited = {origin}
        while cursor is not None:
            if cursor in visited:
                issues.append(_official_row_issue(
                    spec, row_number, "CYCLIC_REFERENCE",
                    "Category parent references contain a cycle.",
                    field="parent_category_code", value=values.get("parent_category_code"),
                ))
                break
            visited.add(cursor)
            cursor = category_parents.get(cursor)
    section_menus = {
        values.get("section_code"): values.get("menu_code")
        for spec, _, values, _ in rows if spec.name == "Menu_Sections"
    }
    for spec, row_number, values, issues in rows:
        if spec.name == "Menu_Items" and values.get("section_code") in section_menus:
            if section_menus[values["section_code"]] != values.get("menu_code"):
                issues.append(_official_row_issue(
                    spec, row_number, "REFERENCE_SCOPE_MISMATCH",
                    "section_code must belong to the referenced menu_code.",
                    field="section_code", value=values.get("section_code"),
                ))


def _official_row_locations(
    rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]],
    indexes: dict[
        tuple[str, str],
        dict[Any, list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]]],
    ],
) -> dict[tuple[str, int], Any]:
    locations: dict[tuple[str, int], Any] = {}
    for spec, row_number, values, _ in rows:
        if values.get("location_code") is not None:
            locations[(spec.name, row_number)] = values["location_code"]
    menu_locations = {
        values.get("menu_code"): values.get("location_code")
        for spec, _, values, _ in rows
        if spec.name == "Menus" and values.get("menu_code") is not None
    }
    warehouse_locations = {
        values.get("warehouse_code"): values.get("location_code")
        for spec, _, values, _ in rows
        if spec.name == "Warehouses" and values.get("warehouse_code") is not None
    }
    inventory_locations = {
        values.get("inventory_code"): warehouse_locations.get(values.get("warehouse_code"))
        for spec, _, values, _ in rows
        if spec.name == "Inventory_Items" and values.get("inventory_code") is not None
    }
    for spec, row_number, values, _ in rows:
        key = (spec.name, row_number)
        if key in locations:
            continue
        if values.get("menu_code") is not None:
            locations[key] = menu_locations.get(values.get("menu_code"))
        elif values.get("warehouse_code") is not None:
            locations[key] = warehouse_locations.get(values.get("warehouse_code"))
        elif values.get("inventory_code") is not None:
            locations[key] = inventory_locations.get(values.get("inventory_code"))
    return locations


def _validate_official_dataset(
    rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]],
    metadata: dict[str, Any],
) -> None:
    confirmed_organizations = [row for row in rows if row[0].name == "Organizations" and row[2].get("capture_status") == "CONFIRMADO"]
    if metadata.get("capture_status") == "CONFIRMADO":
        if len(confirmed_organizations) != 1:
            issue = _workbook_issue(
                "ORGANIZATION_CARDINALITY",
                "Confirmed Metadata requires exactly one confirmed Organization row.",
                sheet="Organizations",
            )
            if confirmed_organizations:
                for row in confirmed_organizations:
                    row[3].append(issue)
        elif confirmed_organizations[0][2].get("organization_code") != metadata.get("organization_code"):
            confirmed_organizations[0][3].append(_official_row_issue(
                confirmed_organizations[0][0], confirmed_organizations[0][1],
                "METADATA_SCOPE_MISMATCH",
                "Organization code conflicts with Metadata.organization_code.",
                field="organization_code",
                value=confirmed_organizations[0][2].get("organization_code"),
            ))
        confirmed_locations = [row for row in rows if row[0].name == "Locations" and row[2].get("capture_status") == "CONFIRMADO" and row[2].get("status") == "ACTIVE"]
        headquarters = [row for row in confirmed_locations if row[2].get("is_headquarters") == "YES"]
        if confirmed_organizations and confirmed_organizations[0][2].get("status") == "ACTIVE" and len(headquarters) != 1:
            target_rows = confirmed_locations or confirmed_organizations
            for spec, row_number, _, issues in target_rows:
                issues.append(_official_row_issue(
                    spec, row_number, "INVALID_HEADQUARTERS_CARDINALITY",
                    "An active confirmed Organization requires exactly one confirmed active headquarters Location.",
                    field="is_headquarters" if spec.name == "Locations" else None,
                ))


def _build_official_result(
    version: str,
    working_rows: list[tuple[OfficialSheet, int, dict[str, Any], list[AnalysisIssue]]],
    warnings: list[AnalysisIssue],
    metadata: dict[str, Any],
    *,
    global_errors: list[AnalysisIssue],
) -> AnalysisResult:
    rows = tuple(
        AnalyzedRow(
            group=spec.key,
            sheet=spec.name,
            row=row_number,
            business_key={field: values.get(field) for field in spec.business_key},
            values=values,
            errors=tuple(sorted(issues, key=_issue_sort_key)),
        )
        for spec, row_number, values, issues in working_rows
    )
    all_errors = tuple(sorted((
        *global_errors,
        *(error for row in rows for error in row.errors),
    ), key=_issue_sort_key))
    groups = tuple(
        GroupSummary(
            spec.key, spec.name,
            sum(row.group == spec.key for row in rows),
            sum(row.group == spec.key and row.valid for row in rows),
            sum(row.group == spec.key and not row.valid for row in rows),
        )
        for spec in OFFICIAL_SHEETS
    )
    valid_count = sum(row.valid for row in rows)
    return AnalysisResult(
        contract_version=version,
        status="VALID" if not all_errors else "INVALID",
        total_populated_rows=len(rows), valid_rows=valid_count,
        invalid_rows=len(rows) - valid_count, warning_count=len(warnings),
        persisted_rows=0, groups=groups, rows=rows, errors=all_errors,
        warnings=tuple(sorted(warnings, key=_issue_sort_key)),
        profile=OFFICIAL_PROFILE, metadata=metadata,
    )


def _official_row_issue(
    spec: OfficialSheet,
    row: int,
    code: str,
    message: str,
    *,
    field: str | None = None,
    value: Any = None,
) -> AnalysisIssue:
    return _workbook_issue(
        code, message, group=spec.key, sheet=spec.name, row=row, field=field, value=value,
    )


def _validate_archive_boundary(content: bytes) -> AnalysisIssue | None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > MAX_XLSX_ARCHIVE_MEMBERS:
                return _workbook_issue(
                    "XLSX_ARCHIVE_LIMIT",
                    f"The workbook archive exceeds {MAX_XLSX_ARCHIVE_MEMBERS} members.",
                )
            if sum(entry.file_size for entry in entries) > MAX_XLSX_EXPANDED_BYTES:
                return _workbook_issue(
                    "XLSX_ARCHIVE_LIMIT",
                    f"The expanded workbook exceeds {MAX_XLSX_EXPANDED_BYTES} bytes.",
                )
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                return _workbook_issue("MACROS_NOT_ALLOWED", "Macro-enabled workbooks are not supported.")
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                return _workbook_issue("INVALID_XLSX", "The archive is missing required OOXML workbook parts.")
    except (BadZipFile, OSError):
        return _workbook_issue("INVALID_XLSX", "The XLSX/OOXML archive is malformed.")
    return None


def _validate_structure(workbook: Any) -> tuple[str | None, list[AnalysisIssue], list[AnalysisIssue]]:
    errors: list[AnalysisIssue] = []
    warnings: list[AnalysisIssue] = []
    required_sheets = (*_SYSTEM_SHEETS, *(group.sheet for group in ONBOARDING_CONTRACT))
    missing = [name for name in required_sheets if name not in workbook.sheetnames]
    for sheet in missing:
        errors.append(_workbook_issue("MISSING_SHEET", f'Required sheet "{sheet}" is missing.', sheet=sheet))

    unknown = [name for name in workbook.sheetnames if name not in required_sheets]
    for sheet in sorted(unknown):
        warnings.append(
            _workbook_issue(
                "IGNORED_SHEET",
                f'Additional sheet "{sheet}" is not part of the contract and was ignored.',
                sheet=sheet,
                severity="WARNING",
            )
        )

    version: str | None = None
    if "Metadata" in workbook.sheetnames:
        metadata_sheet = workbook["Metadata"]
        metadata_headers = _trimmed_headers(metadata_sheet)
        if metadata_headers != ("key", "value"):
            errors.append(
                _workbook_issue(
                    "INVALID_METADATA_HEADERS",
                    'Metadata headers must be exactly "key" and "value".',
                    sheet="Metadata",
                    row=1,
                )
            )
        metadata: dict[str, Any] = {}
        duplicate_keys: set[str] = set()
        for row_number in _populated_row_numbers(metadata_sheet):
            row = tuple(
                metadata_sheet.cell(row_number, column).value
                for column in range(1, metadata_sheet.max_column + 1)
            )
            key = _blank_to_none(row[0] if row else None)
            value = _blank_to_none(row[1] if len(row) > 1 else None)
            if key is None and value is None:
                continue
            if not isinstance(key, str):
                errors.append(_workbook_issue("INVALID_METADATA", "Metadata keys must be text.", sheet="Metadata"))
                continue
            key = key.strip()
            if key in metadata:
                duplicate_keys.add(key)
            else:
                metadata[key] = value.strip() if isinstance(value, str) else value
        for key in sorted(duplicate_keys):
            errors.append(
                _workbook_issue("DUPLICATE_METADATA_KEY", f'Duplicate Metadata key "{key}".', sheet="Metadata")
            )
        raw_version = metadata.get("contract_version")
        version = str(raw_version) if raw_version is not None else None
        if version is None:
            errors.append(
                _workbook_issue("MISSING_CONTRACT_VERSION", "Metadata.contract_version is required.", sheet="Metadata")
            )
        elif version != CONTRACT_VERSION:
            errors.append(
                _workbook_issue(
                    "UNSUPPORTED_CONTRACT_VERSION",
                    f'Contract version "{version}" is not supported; expected "{CONTRACT_VERSION}".',
                    sheet="Metadata",
                    value=version,
                )
            )
        expected_metadata = {
            "template_kind": TEMPLATE_KIND,
            "data_rows_start": "2",
            "max_capture_rows_per_sheet": str(MAX_CAPTURE_ROWS),
            "imports_data": "NO",
            "uses_pos": "NO",
        }
        for key, expected in expected_metadata.items():
            actual = metadata.get(key)
            normalized_actual = str(actual) if actual is not None else None
            if normalized_actual != expected:
                errors.append(
                    _workbook_issue(
                        "INVALID_TEMPLATE_METADATA",
                        f'Metadata.{key} must be "{expected}".',
                        sheet="Metadata",
                        field=key,
                        value=actual,
                    )
                )

    if workbook.properties.title != "ECIP Stage 0 Restaurant Onboarding":
        errors.append(
            _workbook_issue(
                "INVALID_TEMPLATE_IDENTITY",
                "Workbook title does not identify the official template.",
            )
        )
    if workbook.properties.subject != CONTRACT_VERSION:
        errors.append(
            _workbook_issue(
                "INVALID_TEMPLATE_IDENTITY",
                f'Workbook subject must be "{CONTRACT_VERSION}".',
                value=workbook.properties.subject,
            )
        )

    for group in ONBOARDING_CONTRACT:
        if group.sheet not in workbook.sheetnames:
            continue
        sheet = workbook[group.sheet]
        if sheet.merged_cells.ranges:
            errors.append(
                _workbook_issue(
                    "MERGED_CELLS_NOT_ALLOWED",
                    "Merged cells are not supported in contract capture sheets.",
                    group=group.key,
                    sheet=group.sheet,
                )
            )
        headers = _trimmed_headers(sheet)
        expected = tuple(field.key for field in group.fields)
        legacy = _LEGACY_HEADERS.get(group.key)
        accepted_legacy = headers == legacy
        duplicates = sorted({header for header in headers if header is not None and headers.count(header) > 1})
        for header in duplicates:
            errors.append(
                _workbook_issue(
                    "DUPLICATE_HEADER",
                    f'Duplicate header "{header}".',
                    group=group.key,
                    sheet=group.sheet,
                    row=1,
                    field=str(header),
                )
            )
        missing_headers = [] if accepted_legacy else [
            header for header in expected if header not in headers
        ]
        for header in missing_headers:
            errors.append(
                _workbook_issue(
                    "MISSING_HEADER",
                    f'Required contract column "{header}" is missing.',
                    group=group.key,
                    sheet=group.sheet,
                    row=1,
                    field=header,
                )
            )
        if headers != expected and not accepted_legacy and not missing_headers and not duplicates:
            errors.append(
                _workbook_issue(
                    "INVALID_HEADER_CONTRACT",
                    "Columns must match the contract exactly and in contract order.",
                    group=group.key,
                    sheet=group.sheet,
                    row=1,
                )
            )
        contract_width = len(headers) if accepted_legacy else len(expected)
        if sheet.max_column > contract_width:
            for row_number in _populated_row_numbers(sheet):
                if any(
                    sheet.cell(row_number, column).data_type == "f"
                    or _blank_to_none(sheet.cell(row_number, column).value) is not None
                    for column in range(contract_width + 1, sheet.max_column + 1)
                ):
                    errors.append(
                        _workbook_issue(
                            "UNSUPPORTED_COLUMN_DATA",
                            "Data exists outside the contract columns.",
                            group=group.key,
                            sheet=group.sheet,
                            row=row_number,
                        )
                    )
    return version, errors, warnings


def _trimmed_headers(sheet: Any) -> tuple[Any, ...]:
    values = [cell.value.strip() if isinstance(cell.value, str) else cell.value for cell in sheet[1]]
    while values and values[-1] is None:
        values.pop()
    return tuple(values)


def _parse_rows(workbook: Any) -> list[_WorkingRow]:
    rows: list[_WorkingRow] = []
    for group in ONBOARDING_CONTRACT:
        sheet = workbook[group.sheet]
        headers = _trimmed_headers(sheet)
        legacy = headers == _LEGACY_HEADERS.get(group.key)
        if legacy and group.key == "categories":
            fields = _LEGACY_CATEGORY_FIELDS
        elif legacy and group.key == "products":
            fields = _LEGACY_PRODUCT_FIELDS
        else:
            fields = group.fields
        width = len(fields)
        for excel_row in _populated_row_numbers(sheet):
            cells = tuple(sheet.cell(excel_row, column) for column in range(1, width + 1))
            errors: list[AnalysisIssue] = []
            if excel_row > MAX_CAPTURE_ROWS + 1:
                errors.append(
                    _row_issue(
                        group,
                        excel_row,
                        "ROW_LIMIT_EXCEEDED",
                        f"Rows beyond {MAX_CAPTURE_ROWS + 1} are not supported.",
                    )
                )
            values: dict[str, Any] = {}
            for field, cell in zip(fields, cells):
                normalized, error = _normalize_cell(group, excel_row, field, cell)
                values[field.key] = normalized
                if error is not None:
                    errors.append(error)
            if legacy and group.key == "categories":
                category_name = values.pop("category_name")
                parent_name = values.pop("parent_category_name")
                values = {
                    "category_key": category_name,
                    "organization_code": values["organization_code"],
                    "name": category_name,
                    "parent_category_key": parent_name,
                    "display_order": values["display_order"],
                    "status": values["status"],
                }
            elif legacy and group.key == "products":
                values["category_key"] = values.pop("category_name")
            _validate_row_semantics(group, excel_row, values, errors)
            rows.append(_WorkingRow(group, excel_row, values, errors))
    return rows


def _populated_row_numbers(sheet: Any) -> tuple[int, ...]:
    """Find actual data rows without walking large ranges created only by formatting."""

    return tuple(
        sorted(
            {
                cell.row
                for cell in sheet._cells.values()
                if cell.row >= 2 and (cell.data_type == "f" or _blank_to_none(cell.value) is not None)
            }
        )
    )


def _normalize_cell(
    group: Group, excel_row: int, field: Field, cell: Cell
) -> tuple[Any, AnalysisIssue | None]:
    if cell.data_type == "f":
        return None, _row_issue(
            group, excel_row, "FORMULA_NOT_ALLOWED", "Formulas are not authoritative onboarding data.", field=field.key
        )
    raw = _blank_to_none(cell.value)
    if raw is None:
        if field.required:
            return None, _row_issue(
                group, excel_row, "REQUIRED_FIELD", "A required value is missing.", field=field.key
            )
        return None, None
    try:
        if field.value_type in {"text", "enum"}:
            if not isinstance(raw, str):
                raise ValueError("must be text")
            value: Any = raw.strip()
        elif field.value_type == "integer":
            value = _parse_integer(raw)
        elif field.value_type == "decimal":
            value = _parse_decimal(raw)
        elif field.value_type == "datetime":
            value = _parse_datetime(raw)
        else:  # protected by contract validation
            raise ValueError("has an unsupported type")
        _validate_field_value(field, value)
        return value, None
    except (InvalidOperation, TypeError, ValueError) as exc:
        return None, _row_issue(
            group,
            excel_row,
            "INVALID_FIELD_VALUE",
            f"Value {exc} for contract type {field.value_type}.",
            field=field.key,
            value=raw,
        )


def _parse_integer(raw: Any) -> int:
    if isinstance(raw, bool):
        raise ValueError("must be an integer")
    if isinstance(raw, int):
        return raw
    if isinstance(raw, float) and raw.is_integer():
        return int(raw)
    if isinstance(raw, str) and re.fullmatch(r"[+-]?\d+", raw.strip()):
        return int(raw.strip())
    raise ValueError("must be an integer")


def _parse_decimal(raw: Any) -> str:
    if isinstance(raw, bool):
        raise ValueError("must be a decimal")
    value = Decimal(str(raw).strip())
    if not value.is_finite():
        raise ValueError("must be a finite decimal")
    return format(value, "f")


def _parse_datetime(raw: Any) -> str:
    if isinstance(raw, datetime):
        value = raw
    elif isinstance(raw, str):
        text = raw.strip()
        value = datetime.fromisoformat(text[:-1] + "+00:00" if text.endswith("Z") else text)
    else:
        raise ValueError("must be an ISO 8601 datetime")
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("must include a timezone offset")
    return value.isoformat()


def _validate_field_value(field: Field, value: Any) -> None:
    if field.allowed_values and value not in field.allowed_values:
        raise ValueError(f'must be one of {", ".join(field.allowed_values)}')
    fmt = field.format or ""
    if field.value_type == "text":
        if fmt == "slug" and not _SLUG_PATTERN.fullmatch(value):
            raise ValueError("must be a lowercase URL-safe slug")
        if fmt.startswith("code:"):
            maximum = int(fmt.partition(":")[2])
            if len(value) > maximum or not _CODE_PATTERN.fullmatch(value):
                raise ValueError(f"must be an uppercase operator code of at most {maximum} characters")
        if fmt == "email" and (len(value) > 320 or not _EMAIL_PATTERN.fullmatch(value)):
            raise ValueError("must be a valid email address")
        if fmt == "IANA timezone":
            try:
                ZoneInfo(value)
            except (ValueError, ZoneInfoNotFoundError):
                raise ValueError("must be a valid IANA timezone") from None
        if fmt == "BCP-47" and not _LANGUAGE_PATTERN.fullmatch(value):
            raise ValueError("must be a BCP-47-style language tag")
        if fmt == "unit code:32" and not _UNIT_PATTERN.fullmatch(value):
            raise ValueError("must be an uppercase unit code of at most 32 characters")
        if fmt.startswith("^") and not re.fullmatch(fmt, value):
            raise ValueError(f"does not match {fmt}")
    if field.value_type == "integer":
        if ">=0" in fmt and value < 0:
            raise ValueError("must be zero or greater")
        if ">0" in fmt and value <= 0:
            raise ValueError("must be greater than zero")
    if field.value_type == "decimal":
        decimal_value = Decimal(value)
        compact_format = fmt.replace(" ", "")
        shape = _DECIMAL_SHAPE.search(fmt)
        if shape:
            precision, scale = (int(part) for part in shape.groups())
            actual_scale = max(0, -decimal_value.as_tuple().exponent)
            integer_digits = max(0, decimal_value.adjusted() + 1) if decimal_value else 0
            if actual_scale > scale or integer_digits > precision - scale:
                raise ValueError(f"must fit decimal({precision},{scale})")
        if ">=0" in compact_format and decimal_value < 0:
            raise ValueError("must be zero or greater")
        if ">0" in compact_format and decimal_value <= 0:
            raise ValueError("must be greater than zero")


def _validate_row_semantics(
    group: Group, row: int, values: dict[str, Any], errors: list[AnalysisIssue]
) -> None:
    minimum = values.get("min_selections")
    maximum = values.get("max_selections")
    if group.key == "choice_groups" and minimum is not None and maximum is not None and maximum < minimum:
        errors.append(
            _row_issue(
                group,
                row,
                "INVALID_RANGE",
                "max_selections must be greater than or equal to min_selections.",
                field="max_selections",
            )
        )
    if group.key == "promotions":
        benefit = values.get("benefit_value")
        if (
            values.get("promotion_type") == "PERCENTAGE_DISCOUNT"
            and benefit is not None
            and Decimal(benefit) > Decimal("100")
        ):
            errors.append(
                _row_issue(
                    group,
                    row,
                    "INVALID_RANGE",
                    "Percentage benefit_value cannot exceed 100.",
                    field="benefit_value",
                )
            )
        if values.get("promotion_type") == "FIXED_AMOUNT_DISCOUNT" and values.get("currency") is None:
            errors.append(
                _row_issue(
                    group,
                    row,
                    "CONDITIONALLY_REQUIRED",
                    "currency is required for a fixed amount discount.",
                    field="currency",
                )
            )
    if group.key == "preparation_routes":
        area = values.get("preparation_area_code")
        if values.get("policy") == "AREA" and area is None:
            errors.append(
                _row_issue(
                    group,
                    row,
                    "CONDITIONALLY_REQUIRED",
                    "preparation_area_code is required when policy is AREA.",
                    field="preparation_area_code",
                )
            )
        if values.get("policy") in {"COMPONENTS", "NO_PREPARATION"} and area is not None:
            errors.append(
                _row_issue(
                    group,
                    row,
                    "FIELD_NOT_ALLOWED",
                    "preparation_area_code is only allowed when policy is AREA.",
                    field="preparation_area_code",
                    value=area,
                )
            )
    if (
        group.key == "payment_methods"
        and values.get("enabled") == "YES"
        and values.get("method") in {"CARD", "TRANSFER"}
        and values.get("provider") is None
    ):
        errors.append(
            _row_issue(
                group,
                row,
                "CONDITIONALLY_REQUIRED",
                "provider is required for an enabled non-cash method.",
                field="provider",
            )
        )
    if group.key in {"promotions", "preparation_recipes", "tax_rules", "product_fiscal_classifications"}:
        start_key = "starts_at" if group.key == "promotions" else "effective_from"
        end_key = "ends_at" if group.key == "promotions" else "effective_to"
        start = values.get(start_key)
        end = values.get(end_key)
        if start is not None and end is not None and datetime.fromisoformat(end) <= datetime.fromisoformat(start):
            errors.append(
                _row_issue(
                    group,
                    row,
                    "INVALID_DATE_RANGE",
                    f"{end_key} must be later than {start_key}.",
                    field=end_key,
                )
            )


def _validate_duplicates(rows: list[_WorkingRow]) -> None:
    by_group: dict[str, dict[tuple[Any, ...], list[_WorkingRow]]] = {}
    for row in rows:
        key = tuple(row.values.get(field) for field in row.group.business_key)
        if any(value is None for value in key):
            continue
        by_group.setdefault(row.group.key, {}).setdefault(key, []).append(row)
    for groups in by_group.values():
        for key, duplicates in groups.items():
            if len(duplicates) < 2:
                continue
            display = " + ".join(str(value) for value in key)
            for row in duplicates:
                row.errors.append(
                    _row_issue(
                        row.group,
                        row.row,
                        "DUPLICATE_BUSINESS_KEY",
                        f'Duplicate business key "{display}" in this workbook.',
                        field=" + ".join(row.group.business_key),
                        value=display,
                    )
                )


def _validate_references(rows: list[_WorkingRow]) -> dict[int, dict[str, Any]]:
    rows_by_group: dict[str, list[_WorkingRow]] = {group.key: [] for group in ONBOARDING_CONTRACT}
    for row in rows:
        rows_by_group[row.group.key].append(row)
    groups = {group.key: group for group in ONBOARDING_CONTRACT}
    contexts: dict[int, dict[str, Any]] = {}
    for row in rows:
        context = dict(row.values)
        contexts[id(row)] = context
        for field in row.group.fields:
            if not field.reference or row.values.get(field.key) is None:
                continue
            target_group_key, target_field = field.reference.split(".", 1)
            target_group = groups[target_group_key]
            candidates = [
                target
                for target in rows_by_group[target_group_key]
                if target.values.get(target_field) == row.values[field.key]
            ]
            shared_scope = [
                key
                for key in target_group.business_key
                if key != target_field and context.get(key) is not None
            ]
            if shared_scope:
                candidates = [
                    target
                    for target in candidates
                    if all(target.values.get(key) == context.get(key) for key in shared_scope)
                ]
            if len(candidates) == 1:
                target_context = contexts.get(id(candidates[0]), candidates[0].values)
                for key, value in target_context.items():
                    if value is not None and context.get(key) is None:
                        context[key] = value
                continue
            code = "UNRESOLVED_REFERENCE" if not candidates else "AMBIGUOUS_REFERENCE"
            description = "does not resolve" if not candidates else "resolves to multiple rows"
            row.errors.append(
                _row_issue(
                    row.group,
                    row.row,
                    code,
                    f'{field.key} {description} in {target_group.sheet}.{target_field}.',
                    field=field.key,
                    value=row.values[field.key],
                )
            )
    return contexts


def _validate_dataset_rules(
    rows: list[_WorkingRow], row_contexts: dict[int, dict[str, Any]]
) -> None:
    by_group: dict[str, list[_WorkingRow]] = {group.key: [] for group in ONBOARDING_CONTRACT}
    for row in rows:
        by_group[row.group.key].append(row)

    defaults: dict[Any, list[_WorkingRow]] = {}
    for row in by_group["warehouses"]:
        if row.values.get("is_default") == "YES" and row.values.get("location_code") is not None:
            defaults.setdefault(row.values.get("location_code"), []).append(row)
    for duplicate_defaults in defaults.values():
        if len(duplicate_defaults) > 1:
            for row in duplicate_defaults:
                row.errors.append(
                    _row_issue(
                        row.group,
                        row.row,
                        "MULTIPLE_DEFAULTS",
                        "Only one default warehouse is allowed per location.",
                        field="is_default",
                    )
                )

    yield_rows: dict[Any, list[_WorkingRow]] = {}
    for row in by_group["preparation_recipe_components"]:
        if row.values.get("yield_basis") == "YES" and row.values.get("recipe_key") is not None:
            yield_rows.setdefault(row.values.get("recipe_key"), []).append(row)
    for duplicates in yield_rows.values():
        if len(duplicates) > 1:
            for row in duplicates:
                row.errors.append(
                    _row_issue(
                        row.group,
                        row.row,
                        "MULTIPLE_YIELD_BASES",
                        "Only one component per recipe may have yield_basis YES.",
                        field="yield_basis",
                    )
                )

    components = {
        (row.values.get("product_key"), row.values.get("location_code"))
        for row in by_group["consumption_components"]
    }
    for row in by_group["consumption_definitions"]:
        key = (row.values.get("product_key"), row.values.get("location_code"))
        if (
            row.values.get("tracking_mode") == "DERIVABLE"
            and all(value is not None for value in key)
            and key not in components
        ):
            row.errors.append(
                _row_issue(
                    row.group,
                    row.row,
                    "MISSING_DEPENDENT_ROW",
                    "DERIVABLE consumption requires at least one consumption component.",
                )
            )

    sections_by_key = {
        row.values.get("section_key"): row.values.get("menu_key")
        for row in by_group["menu_sections"]
        if row.values.get("section_key") is not None
    }
    for row in by_group["menu_items"]:
        section_key = row.values.get("section_key")
        menu_key = row.values.get("menu_key")
        section_menu_key = sections_by_key.get(section_key)
        if section_menu_key is not None and menu_key is not None and section_menu_key != menu_key:
            row.errors.append(
                _row_issue(
                    row.group,
                    row.row,
                    "REFERENCE_SCOPE_MISMATCH",
                    "section_key must belong to the referenced menu_key.",
                    field="section_key",
                    value=section_key,
                )
            )

    _validate_uom_usage(rows, by_group, row_contexts)


def _validate_uom_usage(
    rows: list[_WorkingRow],
    by_group: dict[str, list[_WorkingRow]],
    row_contexts: dict[int, dict[str, Any]],
) -> None:
    item_uoms: dict[tuple[Any, Any], set[Any]] = {}
    for row in by_group["inventory_items"]:
        key = (row.values.get("location_code"), row.values.get("inventory_item_code"))
        item_uoms.setdefault(key, set()).add(row.values.get("base_uom"))
    for row in by_group["uom_conversions"]:
        key = (row.values.get("location_code"), row.values.get("inventory_item_code"))
        item_uoms.setdefault(key, set()).add(row.values.get("operational_uom"))
    uses = {
        "supplier_offerings": ("purchase_uom", "inventory_item_code"),
        "consumption_components": ("uom", "inventory_item_code"),
        "preparation_recipes": ("output_uom", "output_inventory_item_code"),
        "preparation_recipe_components": ("source_uom", "inventory_item_code"),
    }
    for row in rows:
        if row.group.key not in uses:
            continue
        uom_field, item_field = uses[row.group.key]
        uom = row.values.get(uom_field)
        if uom is None or row.values.get(item_field) is None:
            continue
        location_code = row_contexts[id(row)].get("location_code")
        if location_code is None:
            continue
        key = (location_code, row.values.get(item_field))
        if uom not in item_uoms.get(key, set()):
            row.errors.append(
                _row_issue(
                    row.group,
                    row.row,
                    "UNRESOLVED_UOM",
                    "Unit must be the item's base UOM or a declared operational conversion.",
                    field=uom_field,
                    value=uom,
                )
            )


def _build_result(
    version: str | None, working_rows: list[_WorkingRow], warnings: list[AnalysisIssue]
) -> AnalysisResult:
    rows: list[AnalyzedRow] = []
    for item in working_rows:
        errors = tuple(sorted(item.errors, key=_issue_sort_key))
        rows.append(
            AnalyzedRow(
                group=item.group.key,
                sheet=item.group.sheet,
                row=item.row,
                business_key={key: item.values.get(key) for key in item.group.business_key},
                values=item.values,
                errors=errors,
            )
        )
    rows.sort(key=lambda item: (_group_order(item.group), item.row))
    all_errors = tuple(sorted((error for row in rows for error in row.errors), key=_issue_sort_key))
    summaries = tuple(
        GroupSummary(
            group=group.key,
            sheet=group.sheet,
            populated_rows=sum(row.group == group.key for row in rows),
            valid_rows=sum(row.group == group.key and row.valid for row in rows),
            invalid_rows=sum(row.group == group.key and not row.valid for row in rows),
        )
        for group in ONBOARDING_CONTRACT
    )
    valid_count = sum(row.valid for row in rows)
    return AnalysisResult(
        contract_version=version,
        status="VALID" if not all_errors else "INVALID",
        total_populated_rows=len(rows),
        valid_rows=valid_count,
        invalid_rows=len(rows) - valid_count,
        warning_count=len(warnings),
        persisted_rows=0,
        groups=summaries,
        rows=tuple(rows),
        errors=all_errors,
        warnings=tuple(sorted(warnings, key=_issue_sort_key)),
    )


def _empty_result(
    version: str | None,
    errors: list[AnalysisIssue],
    warnings: list[AnalysisIssue] | None = None,
    *,
    profile: str = "legacy-generated",
    metadata: dict[str, Any] | None = None,
    groups: tuple[GroupSummary, ...] | None = None,
) -> AnalysisResult:
    warnings = warnings or []
    return AnalysisResult(
        contract_version=version,
        status="INVALID" if errors else "VALID",
        total_populated_rows=0,
        valid_rows=0,
        invalid_rows=0,
        warning_count=len(warnings),
        persisted_rows=0,
        groups=(
            groups
            if groups is not None
            else tuple(GroupSummary(group.key, group.sheet, 0, 0, 0) for group in ONBOARDING_CONTRACT)
        ),
        rows=(),
        errors=tuple(sorted(errors, key=_issue_sort_key)),
        warnings=tuple(sorted(warnings, key=_issue_sort_key)),
        profile=profile,
        metadata=metadata,
    )


def _workbook_issue(
    code: str,
    message: str,
    *,
    group: str | None = None,
    sheet: str | None = None,
    row: int | None = None,
    field: str | None = None,
    value: Any = None,
    severity: str = "ERROR",
) -> AnalysisIssue:
    return AnalysisIssue(code, message, group, sheet, row, field, _safe_value(value), severity)


def _row_issue(
    group: Group,
    row: int,
    code: str,
    message: str,
    *,
    field: str | None = None,
    value: Any = None,
) -> AnalysisIssue:
    return _workbook_issue(
        code, message, group=group.key, sheet=group.sheet, row=row, field=field, value=value
    )


def _safe_value(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= 120 else text[:117] + "..."


def _blank_to_none(value: Any) -> Any:
    return None if value is None or (isinstance(value, str) and not value.strip()) else value


def _group_order(group_key: str | None) -> int:
    if group_key is None:
        return -1
    legacy = next((group.order for group in ONBOARDING_CONTRACT if group.key == group_key), None)
    if legacy is not None:
        return legacy
    return next(
        (index for index, group in enumerate(OFFICIAL_SHEETS, start=1) if group.key == group_key),
        10**9,
    )


def _issue_sort_key(issue: AnalysisIssue) -> tuple[Any, ...]:
    return (_group_order(issue.group), issue.sheet or "", issue.row or 0, issue.field or "", issue.code)


__all__ = [
    "AnalysisIssue",
    "AnalysisResult",
    "AnalyzedRow",
    "GroupSummary",
    "MAX_XLSX_ARCHIVE_MEMBERS",
    "MAX_XLSX_BYTES",
    "MAX_XLSX_EXPANDED_BYTES",
    "analyze_xlsx",
]
