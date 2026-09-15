"""Deterministic XLSX rendering for the Stage 0 onboarding contract."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from app.onboarding.contract import (
    CONTRACT_VERSION,
    MAX_CAPTURE_ROWS,
    ONBOARDING_CONTRACT,
    Field,
    Group,
    validate_contract,
)


FIXED_DOCUMENT_TIME = datetime(2026, 9, 14, 0, 0, 0)
FIXED_ZIP_TIME = (2020, 1, 1, 0, 0, 0)
FIXED_MODIFIED_XML = b"2026-09-14T00:00:00Z"
HEADER_REQUIRED = PatternFill("solid", fgColor="1F4E78")
HEADER_OPTIONAL = PatternFill("solid", fgColor="5B6573")
HEADER_FONT = Font(color="FFFFFF", bold=True)
META_FILL = PatternFill("solid", fgColor="DDEBF7")
INVALID_FILL = PatternFill("solid", fgColor="FCE4D6")


def _field_comment(field: Field) -> str:
    requirement = "OBLIGATORIO" if field.required else "OPCIONAL"
    details = [requirement, f"Tipo: {field.value_type}", field.description]
    if field.format:
        details.append(f"Formato: {field.format}")
    if field.reference:
        details.append(f"Referencia: {field.reference}")
    if field.allowed_values:
        details.append("Valores: " + ", ".join(field.allowed_values))
    return "\n".join(details)


def _enum_catalog(groups: tuple[Group, ...]) -> tuple[tuple[str, tuple[str, ...]], ...]:
    values = {field.allowed_values for group in groups for field in group.fields if field.allowed_values}
    return tuple((f"enum_{index:02d}", value) for index, value in enumerate(sorted(values), start=1))


def build_workbook(groups: tuple[Group, ...] = ONBOARDING_CONTRACT) -> Workbook:
    validate_contract(groups)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "ECIP"
    workbook.properties.lastModifiedBy = "ECIP"
    workbook.properties.title = "ECIP Stage 0 Restaurant Onboarding"
    workbook.properties.subject = CONTRACT_VERSION
    workbook.properties.description = "Capture-only template; no import or POS integration."
    workbook.properties.created = FIXED_DOCUMENT_TIME
    workbook.properties.modified = FIXED_DOCUMENT_TIME

    instructions = workbook.create_sheet("Instructions")
    instructions.sheet_view.showGridLines = False
    instruction_rows = (
        ("ECIP Stage 0 Restaurant Onboarding",),
        ("Contract version", CONTRACT_VERSION),
        ("Purpose", "Capture configuration for later analyze/validate/preview. This workbook does not import data."),
        ("References", "Use codes and workbook-local keys exactly as written in their source sheet; never enter database IDs."),
        ("Required fields", "Dark blue headers are required; gray headers are optional. Hover headers for constraints."),
        ("Dates", "Use ISO 8601 with timezone where requested, for example 2026-09-14T10:00:00-06:00."),
        ("Decimals", "Enter exact decimal values, never currency symbols or thousands separators."),
        ("Credentials", "Do not enter passwords, hashes, tokens, API keys, access codes or other secrets."),
        ("Inventory history", "Do not enter stock movements, lots, receipts, losses, batches, transfers or valuations."),
        ("Opening stock", "Opening stock must later use an authorized inventory operation; it is not captured here."),
        ("POS boundary", "No POS access or integration is used or required."),
        ("Applicability", "Optional sheets may remain empty when the restaurant confirms that capability is not used."),
        (),
        ("Sheet", "Group key", "Stage 0", "Purpose", "Business key", "Authority"),
    )
    for row in instruction_rows:
        instructions.append(row)
    for group in groups:
        instructions.append((group.sheet, group.key, "REQUIRED" if group.required_for_stage0 else "IF APPLICABLE", group.purpose, " + ".join(group.business_key), group.authority))
    instructions.freeze_panes = "A15"
    instructions.auto_filter.ref = f"A14:F{14 + len(groups)}"
    instructions.column_dimensions["A"].width = 27
    instructions.column_dimensions["B"].width = 34
    instructions.column_dimensions["C"].width = 15
    instructions.column_dimensions["D"].width = 55
    instructions.column_dimensions["E"].width = 42
    instructions.column_dimensions["F"].width = 55
    instructions["A1"].font = Font(size=18, bold=True, color="1F4E78")
    for cell in instructions[14]:
        cell.fill = HEADER_REQUIRED
        cell.font = HEADER_FONT

    metadata = workbook.create_sheet("Metadata")
    metadata.append(("key", "value"))
    metadata.append(("contract_version", CONTRACT_VERSION))
    metadata.append(("template_kind", "stage0-real-restaurant-capture"))
    metadata.append(("data_rows_start", "2"))
    metadata.append(("max_capture_rows_per_sheet", str(MAX_CAPTURE_ROWS)))
    metadata.append(("imports_data", "NO"))
    metadata.append(("uses_pos", "NO"))
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = "A1:B7"
    metadata.column_dimensions["A"].width = 34
    metadata.column_dimensions["B"].width = 48
    for cell in metadata[1]:
        cell.fill = HEADER_REQUIRED
        cell.font = HEADER_FONT

    enum_catalog = _enum_catalog(groups)
    enum_lookup = {values: column for column, (_, values) in enumerate(enum_catalog, start=1)}
    lists = workbook.create_sheet("_Lists")
    for column, (name, values) in enumerate(enum_catalog, start=1):
        lists.cell(1, column, name)
        for row, value in enumerate(values, start=2):
            lists.cell(row, column, value)
    lists.sheet_state = "veryHidden"

    for group in groups:
        sheet = workbook.create_sheet(group.sheet)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for column, field in enumerate(group.fields, start=1):
            cell = sheet.cell(1, column, field.key)
            cell.font = HEADER_FONT
            cell.fill = HEADER_REQUIRED if field.required else HEADER_OPTIONAL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.comment = Comment(_field_comment(field), "ECIP")
            width = max(len(field.key) + 2, min(34, max(14, len(field.label) + 4)))
            sheet.column_dimensions[get_column_letter(column)].width = width
            if field.allowed_values:
                enum_column = enum_lookup[field.allowed_values]
                enum_letter = get_column_letter(enum_column)
                formula = f"{quote_sheetname(lists.title)}!${enum_letter}$2:${enum_letter}${len(field.allowed_values) + 1}"
                validation = DataValidation(type="list", formula1=formula, allow_blank=not field.required)
                validation.error = "Seleccione un valor permitido por el contrato."
                validation.errorTitle = "Valor no permitido"
                validation.prompt = field.description
                validation.promptTitle = field.label
                validation.showErrorMessage = True
                validation.showInputMessage = True
                sheet.add_data_validation(validation)
                validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{MAX_CAPTURE_ROWS + 1}")
            if field.required:
                letter = get_column_letter(column)
                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{MAX_CAPTURE_ROWS + 1}",
                    FormulaRule(formula=[f'AND(COUNTA($A2:${get_column_letter(len(group.fields))}2)>0,{letter}2="")'], fill=INVALID_FILL),
                )
        sheet.row_dimensions[1].height = 36
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(group.fields))}{MAX_CAPTURE_ROWS + 1}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(group.fields))}{MAX_CAPTURE_ROWS + 1}"

    workbook.active = 0
    return workbook


def deterministic_bytes(groups: tuple[Group, ...] = ONBOARDING_CONTRACT) -> bytes:
    raw = BytesIO()
    build_workbook(groups).save(raw)
    source = BytesIO(raw.getvalue())
    target = BytesIO()
    with ZipFile(source, "r") as input_zip, ZipFile(target, "w", compression=ZIP_DEFLATED, compresslevel=9) as output_zip:
        for name in sorted(input_zip.namelist()):
            info = ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 3
            content = input_zip.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(?<=<dcterms:modified xsi:type=\"dcterms:W3CDTF\">).*?(?=</dcterms:modified>)",
                    FIXED_MODIFIED_XML,
                    content,
                )
            output_zip.writestr(info, content, compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()


def generate_template(path: Path, groups: tuple[Group, ...] = ONBOARDING_CONTRACT) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(deterministic_bytes(groups))
    return path
