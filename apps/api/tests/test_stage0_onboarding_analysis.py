from __future__ import annotations

from io import BytesIO
import inspect
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from app.onboarding.analyzer import MAX_XLSX_BYTES, analyze_xlsx
from app.onboarding.contract import CONTRACT_VERSION
from app.onboarding.xlsx import deterministic_bytes


def _edited_workbook(edit) -> bytes:
    workbook = load_workbook(BytesIO(deterministic_bytes()), data_only=False)
    edit(workbook)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _set_row(workbook, sheet: str, values: tuple[object, ...], row: int = 2) -> None:
    for column, value in enumerate(values, start=1):
        workbook[sheet].cell(row, column, value)


def _add_restaurant(workbook) -> None:
    _set_row(
        workbook,
        "01_Restaurant",
        (
            " pilot ",
            " Restaurante Piloto ",
            "ORG-01",
            "Organización",
            "LOC-01",
            "Centro",
            "America/Mexico_City",
            None,
            None,
            None,
            None,
            None,
            "MX",
            None,
            "ops@example.test",
            "ACTIVE",
        ),
    )


def _add_valid_menu_dataset(workbook) -> None:
    _add_restaurant(workbook)
    _set_row(workbook, "06_Categories", ("TACOS", "ORG-01", "Tacos", None, 0, "ACTIVE"))
    _set_row(workbook, "07_Products", ("TACO-1", "ORG-01", "TACOS", " Taco al pastor ", None, "ACTIVE"))
    _set_row(workbook, "08_Menus", ("MENU-1", "ORG-01", "Cena", "LOC-01", "ACTIVE"))
    _set_row(workbook, "09_Menu_Sections", ("SEC-1", "MENU-1", "Principal", " 2 ", "ACTIVE"))
    _set_row(workbook, "10_Menu_Items", ("MENU-1", "SEC-1", "TACO-1", 1, "ACTIVE"))
    _set_row(workbook, "11_Prices", ("TACO-1", "LOC-01", "12.3400", "MXN", "ACTIVE"))


def test_official_empty_workbook_is_recognized_without_business_rows() -> None:
    official = (
        Path(__file__).resolve().parents[3]
        / "docs/operations/templates/ECIP_Stage0_Restaurant_Onboarding_v1.xlsx"
    )

    result = analyze_xlsx(official.read_bytes(), filename=official.name)

    assert result.status == "VALID"
    assert result.contract_version == CONTRACT_VERSION
    assert result.total_populated_rows == result.valid_rows == result.invalid_rows == 0
    assert result.persisted_rows == 0
    assert len(result.groups) == 32
    assert result.to_dict()["persisted_rows"] == 0


def test_valid_rows_normalize_and_resolve_deterministically() -> None:
    content = _edited_workbook(_add_valid_menu_dataset)

    first = analyze_xlsx(content)
    second = analyze_xlsx(content)

    assert first.status == "VALID"
    assert first.total_populated_rows == first.valid_rows == 7
    assert first.invalid_rows == 0
    assert first.to_dict() == second.to_dict()
    restaurant = next(row for row in first.rows if row.group == "restaurant_profile")
    section = next(row for row in first.rows if row.group == "menu_sections")
    price = next(row for row in first.rows if row.group == "prices")
    assert restaurant.values["tenant_slug"] == "pilot"
    assert restaurant.values["tenant_name"] == "Restaurante Piloto"
    assert section.values["display_order"] == 2
    assert price.values["amount"] == "12.3400"
    assert all(not row.errors for row in first.rows)


def test_version_missing_sheet_and_header_contract_fail_before_row_parsing() -> None:
    def edit(workbook) -> None:
        workbook["Metadata"]["B2"] = "restaurant-onboarding/v999"
        workbook.remove(workbook["21_Suppliers"])
        workbook["07_Products"]["B1"] = "product_key"

    result = analyze_xlsx(_edited_workbook(edit))
    codes = {error.code for error in result.errors}

    assert result.status == "INVALID"
    assert result.contract_version == "restaurant-onboarding/v999"
    assert result.total_populated_rows == 0
    assert {"UNSUPPORTED_CONTRACT_VERSION", "MISSING_SHEET", "DUPLICATE_HEADER", "MISSING_HEADER"} <= codes


def test_data_outside_contract_columns_is_rejected_as_structural_ambiguity() -> None:
    def edit(workbook) -> None:
        workbook["07_Products"]["G2"] = "not a contract field"

    result = analyze_xlsx(_edited_workbook(edit))

    assert result.status == "INVALID"
    assert result.total_populated_rows == 0
    assert any(error.code == "UNSUPPORTED_COLUMN_DATA" for error in result.errors)


def test_non_xlsx_corrupt_and_oversized_inputs_are_rejected() -> None:
    malformed_archive = BytesIO()
    with ZipFile(malformed_archive, "w") as archive:
        archive.writestr("unrelated.txt", "not a workbook")
    malformed_ooxml = BytesIO()
    with ZipFile(malformed_ooxml, "w") as archive:
        archive.writestr("[Content_Types].xml", "not xml")
        archive.writestr("xl/workbook.xml", "not xml")
    wrong_extension = analyze_xlsx(deterministic_bytes(), filename="onboarding.xls")
    corrupt = analyze_xlsx(b"not an OOXML workbook", filename="onboarding.xlsx")
    malformed = analyze_xlsx(malformed_archive.getvalue(), filename="onboarding.xlsx")
    malformed_xml = analyze_xlsx(malformed_ooxml.getvalue(), filename="onboarding.xlsx")
    oversized = analyze_xlsx(b"x" * (MAX_XLSX_BYTES + 1), filename="onboarding.xlsx")

    invalid_results = (wrong_extension, corrupt, malformed, malformed_xml, oversized)
    assert [result.status for result in invalid_results] == ["INVALID"] * 5
    assert wrong_extension.errors[0].code == "INVALID_FILE_TYPE"
    assert corrupt.errors[0].code == "INVALID_XLSX"
    assert malformed.errors[0].code == "INVALID_XLSX"
    assert malformed_xml.errors[0].code == "INVALID_XLSX"
    assert oversized.errors[0].code == "FILE_TOO_LARGE"
    assert all(result.persisted_rows == 0 for result in invalid_results)


def test_partial_rows_invalid_values_and_formulas_have_row_provenance() -> None:
    def edit(workbook) -> None:
        workbook["07_Products"]["A2"] = "PROD-1"
        workbook["07_Products"]["D20"] = None
        workbook["07_Products"]["D20"].fill = PatternFill("solid", fgColor="FFFF00")
        _set_row(workbook, "11_Prices", ("UNKNOWN", "LOC-X", "-0.1", "mxn", "ACTIVE"))
        workbook["11_Prices"]["C2"] = "=1+1"

    result = analyze_xlsx(_edited_workbook(edit))
    products = [row for row in result.rows if row.group == "products"]
    price = next(row for row in result.rows if row.group == "prices")

    assert result.status == "INVALID"
    assert len(products) == 1
    assert result.total_populated_rows == 2
    assert {(error.row, error.field, error.code) for error in products[0].errors} >= {
        (2, "organization_code", "REQUIRED_FIELD"),
        (2, "name", "REQUIRED_FIELD"),
        (2, "status", "REQUIRED_FIELD"),
    }
    assert (2, "amount", "FORMULA_NOT_ALLOWED") in {
        (error.row, error.field, error.code) for error in price.errors
    }
    assert all(error.sheet and error.row == 2 for error in price.errors)


def test_invalid_enum_decimal_duplicates_and_unresolved_references_are_reported() -> None:
    def edit(workbook) -> None:
        _add_restaurant(workbook)
        _set_row(workbook, "07_Products", ("DUP-1", "ORG-01", "Missing", "Uno", None, "BROKEN"))
        _set_row(workbook, "07_Products", ("DUP-1", "ORG-01", None, "Dos", None, "ACTIVE"), row=3)
        _set_row(workbook, "11_Prices", ("DUP-1", "LOC-01", "-1.00000", "MXN", "ACTIVE"))

    result = analyze_xlsx(_edited_workbook(edit))
    errors = {(error.group, error.field, error.code) for error in result.errors}

    assert result.status == "INVALID"
    assert result.total_populated_rows == 4
    assert result.valid_rows == 1
    assert result.invalid_rows == 3
    assert ("products", "status", "INVALID_FIELD_VALUE") in errors
    assert ("prices", "amount", "INVALID_FIELD_VALUE") in errors
    assert ("products", "category_key", "UNRESOLVED_REFERENCE") in errors
    assert sum(error.code == "DUPLICATE_BUSINESS_KEY" for error in result.errors) == 2
    assert any(error.code == "AMBIGUOUS_REFERENCE" for error in result.errors)


def test_scoped_references_uom_and_conditional_rules_validate_without_persistence() -> None:
    def edit(workbook) -> None:
        _add_restaurant(workbook)
        _set_row(
            workbook,
            "19_Inventory_Items",
            ("LOC-01", "TORTILLA", "Tortilla", "G", "0.100000", "MXN", "OPTIONAL", "NONE", "ACTIVE"),
        )
        _set_row(workbook, "20_UOM_Conversions", ("LOC-01", "TORTILLA", "PACK", "12", None, None))
        _set_row(
            workbook,
            "25_Recipes",
            ("REC-1", "LOC-01", "TORTILLA", "12", "PACK", None, None, "ACTIVE"),
        )
        _set_row(workbook, "26_Recipe_Lines", ("REC-1", "TORTILLA", "1", "PACK", "YES"))
        _set_row(workbook, "27_Prep_Routes", ("MISSING", "LOC-01", "AREA", None, "ACTIVE"))

    result = analyze_xlsx(_edited_workbook(edit))
    recipe_line = next(row for row in result.rows if row.group == "preparation_recipe_components")
    route = next(row for row in result.rows if row.group == "preparation_routes")

    assert not recipe_line.errors
    assert {error.code for error in route.errors} == {"CONDITIONALLY_REQUIRED", "UNRESOLVED_REFERENCE"}
    assert result.persisted_rows == 0

    import app.onboarding.analyzer as analyzer_module

    source = inspect.getsource(analyzer_module).lower()
    assert "sqlalchemy" not in source
    assert "asyncsession" not in source
    assert ".commit(" not in source
    assert "session.add(" not in source


def test_legacy_category_name_headers_map_to_explicit_stable_references() -> None:
    def edit(workbook) -> None:
        _add_restaurant(workbook)
        legacy_category_headers = (
            "organization_code", "category_name", "parent_category_name",
            "display_order", "status",
        )
        for column, header in enumerate(legacy_category_headers, start=1):
            workbook["06_Categories"].cell(1, column, header)
        workbook["06_Categories"].cell(1, 6).value = None
        workbook["07_Products"].cell(1, 3, "category_name")
        _set_row(workbook, "06_Categories", ("ORG-01", "Tacos", None, 0, "ACTIVE"))
        workbook["06_Categories"].cell(2, 6).value = None
        _set_row(
            workbook, "07_Products",
            ("TACO-1", "ORG-01", "Tacos", "Taco", None, "ACTIVE"),
        )

    result = analyze_xlsx(_edited_workbook(edit))
    category = next(row for row in result.rows if row.group == "categories")
    product = next(row for row in result.rows if row.group == "products")

    assert result.status == "VALID"
    assert category.values["category_key"] == category.values["name"] == "Tacos"
    assert product.values["category_key"] == "Tacos"
    assert result.persisted_rows == 0


def test_menu_item_section_must_belong_to_its_referenced_menu() -> None:
    def edit(workbook) -> None:
        _add_restaurant(workbook)
        _set_row(workbook, "07_Products", ("PROD-1", "ORG-01", None, "Producto", None, "ACTIVE"))
        _set_row(workbook, "08_Menus", ("MENU-1", "ORG-01", "Uno", "LOC-01", "ACTIVE"))
        _set_row(workbook, "08_Menus", ("MENU-2", "ORG-01", "Dos", "LOC-01", "ACTIVE"), row=3)
        _set_row(workbook, "09_Menu_Sections", ("SEC-1", "MENU-1", "Principal", 0, "ACTIVE"))
        _set_row(workbook, "10_Menu_Items", ("MENU-2", "SEC-1", "PROD-1", 0, "ACTIVE"))

    result = analyze_xlsx(_edited_workbook(edit))
    menu_item = next(row for row in result.rows if row.group == "menu_items")

    assert {(error.field, error.code) for error in menu_item.errors} == {
        ("section_key", "REFERENCE_SCOPE_MISMATCH")
    }
    assert result.persisted_rows == 0


def test_unknown_sheet_is_explicitly_ignored_as_a_non_blocking_warning() -> None:
    content = _edited_workbook(lambda workbook: workbook.create_sheet("Operator Notes"))

    result = analyze_xlsx(content)

    assert result.status == "VALID"
    assert result.warning_count == 1
    assert result.warnings[0].code == "IGNORED_SHEET"
