from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import inspect
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from app.onboarding.contract import CONTRACT_VERSION, MAX_CAPTURE_ROWS, ONBOARDING_CONTRACT, validate_contract
from app.onboarding.xlsx import deterministic_bytes, generate_template


EXPECTED_GROUPS = {
    "restaurant_profile", "staff", "resources", "preparation_configuration",
    "preparation_areas", "categories", "products", "product_aliases", "menus", "menu_sections",
    "menu_items", "prices", "product_compositions", "product_components",
    "choice_groups", "choice_options", "promotions", "promotion_products", "promotion_locations",
    "warehouses", "inventory_items", "uom_conversions", "suppliers",
    "supplier_offerings", "consumption_definitions", "consumption_components",
    "preparation_recipes", "preparation_recipe_components", "preparation_routes",
    "tax_rules", "product_fiscal_classifications", "payment_methods",
}
FORBIDDEN_FIELD_KEYS = {
    "password", "password_hash", "access_code", "access_code_digest", "api_key", "secret",
    "stock_movement", "lot_history", "goods_receipt", "inventory_loss", "preparation_batch",
}


def test_contract_is_versioned_valid_complete_and_deterministic() -> None:
    validate_contract()
    assert CONTRACT_VERSION == "restaurant-onboarding/v1"
    assert {group.key for group in ONBOARDING_CONTRACT} == EXPECTED_GROUPS
    assert tuple(group.order for group in ONBOARDING_CONTRACT) == tuple(sorted(group.order for group in ONBOARDING_CONTRACT))
    assert all(len({field.key for field in group.fields}) == len(group.fields) for group in ONBOARDING_CONTRACT)
    assert deterministic_bytes() == deterministic_bytes()


def test_workbook_matches_contract_and_contains_no_capture_data_or_secrets(tmp_path: Path) -> None:
    output = generate_template(tmp_path / "onboarding.xlsx")
    workbook = load_workbook(output, data_only=False)
    expected_sheets = ["Instructions", "Metadata", "_Lists", *(group.sheet for group in ONBOARDING_CONTRACT)]
    assert workbook.sheetnames == expected_sheets
    assert workbook["Metadata"]["B2"].value == CONTRACT_VERSION
    assert workbook["_Lists"].sheet_state == "veryHidden"
    for group in ONBOARDING_CONTRACT:
        sheet = workbook[group.sheet]
        assert tuple(cell.value for cell in sheet[1]) == tuple(field.key for field in group.fields)
        assert sheet.freeze_panes == "A2"
        assert sheet.max_row == 1
        assert sheet.auto_filter.ref.endswith(str(MAX_CAPTURE_ROWS + 1))
    capture_headers = {
        str(cell.value).lower()
        for group in ONBOARDING_CONTRACT
        for cell in workbook[group.sheet][1]
    }
    assert capture_headers.isdisjoint(FORBIDDEN_FIELD_KEYS)


def test_official_artifact_is_current_and_generation_has_no_persistence_path() -> None:
    official = Path(__file__).resolve().parents[3] / "docs/operations/templates/ECIP_Stage0_Restaurant_Onboarding_v1.xlsx"
    assert official.read_bytes() == deterministic_bytes()
    import app.onboarding.contract as contract_module
    import app.onboarding.xlsx as xlsx_module
    source = inspect.getsource(contract_module) + inspect.getsource(xlsx_module)
    assert "sqlalchemy" not in source.lower()
    assert "databasemanager" not in source.lower()
    assert ".commit(" not in source.lower()


def test_xlsx_archive_is_reproducible_and_has_fixed_metadata() -> None:
    first = deterministic_bytes()
    second = deterministic_bytes()
    assert sha256(first).digest() == sha256(second).digest()
    with ZipFile(BytesIO(first)) as archive:
        assert all(item.date_time == (2020, 1, 1, 0, 0, 0) for item in archive.infolist())
        assert "restaurant-onboarding/v1" in archive.read("docProps/core.xml").decode()
