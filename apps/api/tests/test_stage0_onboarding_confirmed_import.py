from __future__ import annotations

import asyncio
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import create_app
from app.onboarding.importer import coverage
from app.onboarding.contract import CONTRACT_VERSION
from app.onboarding.xlsx import deterministic_bytes
from app.restaurant.catalog import category_provisioning
from app.restaurant.catalog import provisioning as product_provisioning
from app.restaurant import (
    menu_item_provisioning,
    menu_provisioning,
    menu_section_provisioning,
    resource_provisioning,
)
from app.restaurant.preparation import (
    area_provisioning,
    configuration_provisioning,
    route_provisioning,
)
from app.restaurant.pricing import provisioning as price_provisioning
from app.restaurant.inventory import service as inventory_service
from app.restaurant.inventory import consumption_provisioning
from app.restaurant.inventory import warehouse_provisioning
from app.restaurant.fiscal_product import provisioning as fiscal_product_provisioning
from app.restaurant.fiscal_product.contracts import FiscalProductClassificationCandidate
from app.restaurant.fiscal_product.service import resolve_fiscal_product_evidence
from app.restaurant.tax import provisioning as tax_rule_provisioning
from app.restaurant.tax.contracts import RestaurantTaxLineCandidate
from app.restaurant.tax.service import resolve_tax_evidence
from test_inventory_recipe_stock_foundation import _headers, _permission, _scope


MEDIA_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def test_all_contract_groups_have_one_explicit_import_classification() -> None:
    values = coverage()
    assert len(values) == 32
    by_group = {value['group']: value for value in values}
    assert len(by_group) == 32
    assert by_group['inventory_items']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['resources']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['preparation_configuration']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['preparation_areas']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['preparation_routes']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['categories']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['products']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['menus']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['menu_sections']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['menu_items']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['prices']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['consumption_definitions']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['consumption_components']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['staff']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['tax_rules']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['product_fiscal_classifications']['classification'] == (
        'IMPORTABLE_NOW'
    )
    assert by_group['warehouses']['classification'] == 'IMPORTABLE_NOW'
    assert by_group['payment_methods']['classification'] == 'OPERATIONAL_NOT_CATALOG_IMPORT'
    assert all(value['authority'] and value['business_key'] for value in values)


def _set_row(workbook, sheet: str, values: tuple[object, ...], row: int = 2) -> None:
    for column, value in enumerate(values, start=1):
        workbook[sheet].cell(row, column, value)


def _workbook(
    tenant_slug: str, *, item_code: str = 'FLOUR', product: bool = False,
    category_product: bool = False, resource: bool = False,
    preparation_configuration: bool = False, preparation_area: bool = False,
    area_resource: bool = False, preparation_route: bool = False,
    menu: bool = False, menu_section: bool = False, menu_item: bool = False,
    price: bool = False,
    inventory: bool = True, warehouse: bool = False,
    consumption: bool = False,
    fiscal: bool = False,
) -> bytes:
    workbook = load_workbook(BytesIO(deterministic_bytes()), data_only=False)
    _set_row(workbook, '01_Restaurant', (
        tenant_slug, 'Inventory Tenant', 'ORG', 'Inventory Organization', 'LOC',
        'Inventory Location', 'America/Mexico_City', None, None, None, None,
        None, 'MX', None, None, 'ACTIVE',
    ))
    if inventory:
        _set_row(workbook, '19_Inventory_Items', (
            'LOC', item_code, f'Item {item_code}', 'G', '0.010000', 'MXN',
            'OPTIONAL', 'NONE', 'ACTIVE',
        ))
        _set_row(workbook, '20_UOM_Conversions', (
            'LOC', item_code, 'BAG', '1000', None, 'approved capture',
        ))
        _set_row(workbook, '21_Suppliers', (
            'ORG', f'SUP-{item_code}', f'Supplier {item_code}', None, 'LOC', 'ACTIVE',
        ))
        _set_row(workbook, '22_Supplier_Offers', (
            f'SUP-{item_code}', 'LOC', item_code, None, 'G', 'ACTIVE',
        ))
    if category_product:
        _set_row(workbook, '06_Categories', (
            f'CAT-{item_code}', 'ORG', f'Category {item_code}', None, 0, 'ACTIVE',
        ))
    if product or category_product or menu_item or price or consumption or fiscal:
        _set_row(workbook, '07_Products', (
            f'PROD-{item_code}', 'ORG',
            f'CAT-{item_code}' if category_product else None,
            f'Product {item_code}', None, 'ACTIVE',
        ))
    if resource:
        _set_row(workbook, '03_Resources', (
            'LOC', f'RESOURCE-{item_code}', f'Resource {item_code}', 'TABLE', 'ACTIVE',
        ))
    if preparation_configuration:
        _set_row(workbook, '04_Prep_Config', ('LOC', 'PLATFORM'))
    if preparation_area:
        _set_row(workbook, '05_Prep_Areas', (
            'LOC', f'AREA-{item_code}', f'Area {item_code}',
            f'RESOURCE-{item_code}' if area_resource else None, 'ACTIVE',
        ))
    if preparation_route:
        _set_row(workbook, '27_Prep_Routes', (
            f'PROD-{item_code}', 'LOC', 'AREA', f'AREA-{item_code}', 'ACTIVE',
        ))
    if menu or menu_section or menu_item:
        _set_row(workbook, '08_Menus', (
            f'MENU-{item_code}', 'ORG', f'Menu {item_code}', 'LOC', 'ACTIVE',
        ))
    if menu_section or menu_item:
        _set_row(workbook, '09_Menu_Sections', (
            f'SECTION-{item_code}', f'MENU-{item_code}',
            f'Section {item_code}', 10, 'ACTIVE',
        ))
    if menu_item:
        _set_row(workbook, '10_Menu_Items', (
            f'MENU-{item_code}', f'SECTION-{item_code}',
            f'PROD-{item_code}', 20, 'ACTIVE',
        ))
    if price:
        _set_row(workbook, '11_Prices', (
            f'PROD-{item_code}', 'LOC', '12.3400', 'MXN', 'ACTIVE',
        ))
    if warehouse:
        _set_row(workbook, '18_Warehouses', (
            'LOC', f'WH-{item_code}', f'Warehouse {item_code}',
            'YES', 'ALLOW', 'ACTIVE',
        ))
    if consumption:
        _set_row(workbook, '23_Consumption', (
            f'PROD-{item_code}', 'LOC', 'DERIVABLE', None, 'ACTIVE',
        ))
        _set_row(workbook, '24_Consumption_Lines', (
            f'PROD-{item_code}', 'LOC', item_code, '2.000000', 'BAG',
        ))
    if fiscal:
        _set_row(workbook, '28_Tax_Rules', (
            f'TAX-{item_code}', 'ORG', None, f'CLASS-{item_code}',
            'MX-FEDERAL', 'IVA', 'TAXABLE', 'TRANSFERRED', '0.160000',
            'INCLUDED_PRICE_SINGLE_TAX', 'DECIMAL_4_HALF_UP',
            '2026-01-01T00:00:00-06:00', None, 'ACTIVE',
        ))
        _set_row(workbook, '29_Product_Fiscal', (
            f'PROD-{item_code}', f'CLASS-{item_code}', 'MX', 'SAT-PRODUCT',
            '50192701', 'SAT-UNIT', 'H87',
            '2026-01-01T00:00:00-06:00', None, 'ACTIVE',
        ))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _prepare(connection, prefix: str):
    scope = _scope(connection, f'{prefix}-onboarding')
    with connection.cursor() as cursor:
        cursor.execute(
            "UPDATE locations SET country_code='MX' WHERE id=%s", (scope.location_id,)
        )
    for permission in (
        'location.manage', 'inventory.manage', 'inventory.supplier.manage', 'product.manage',
        'resource.manage', 'preparation.configure',
        'menu.manage',
        'pricing.manage',
    ):
        _permission(connection, scope.role_id, permission)
    return scope


@pytest.fixture
def client(integration_settings):
    with TestClient(create_app(settings=integration_settings)) as value:
        yield value


def _analyze(client, headers, content):
    return client.post(
        '/onboarding/stage0/analyze', headers={**headers, 'Content-Type': MEDIA_TYPE},
        content=content,
    )


def _confirm(client, headers, content, location_id, fingerprint, *, explicit=True):
    values = {
        **headers, 'Content-Type': MEDIA_TYPE,
        'X-Dataset-Fingerprint': fingerprint,
    }
    if explicit:
        values['X-Confirm-Import'] = 'true'
    return client.post(
        f'/onboarding/stage0/confirm?location_id={location_id}',
        headers=values, content=content,
    )


def _count(connection, table: str, tenant_id: int) -> int:
    with connection.cursor() as cursor:
        cursor.execute(f'SELECT COUNT(*) AS amount FROM {table} WHERE tenant_id=%s', (tenant_id,))
        return int(cursor.fetchone()['amount'])


def test_confirmation_binding_authorities_replay_and_no_inventory_history(client, sql_connection) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding')

    analysis = _analyze(client, headers, content)
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()['analysis']['status'] == 'VALID'
    fingerprint = analysis.json()['dataset_fingerprint']
    assert _count(connection, 'inventory_items', scope.tenant_id) == 0
    assert _count(connection, 'onboarding_imports', scope.tenant_id) == 0

    missing_confirmation = _confirm(
        client, headers, content, scope.location_id, fingerprint, explicit=False,
    )
    assert missing_confirmation.status_code == 422
    changed_content = _workbook(f'{prefix}-onboarding', item_code='CHANGED')
    mismatch = _confirm(client, headers, changed_content, scope.location_id, fingerprint)
    assert mismatch.status_code == 409
    assert _count(connection, 'inventory_items', scope.tenant_id) == 0

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['status'] == 'SUCCESS' and result['replay'] is False, result
    assert 'products' not in result['required_deferred_groups']
    assert (result['total_planned_rows'], result['created'], result['unchanged']) == (5, 4, 1)
    assert [result['groups'][key]['created'] for key in (
        'inventory_items', 'uom_conversions', 'suppliers', 'supplier_offerings'
    )] == [1, 1, 1, 1]
    assert _count(connection, 'inventory_items', scope.tenant_id) == 1
    assert _count(connection, 'stock_movements', scope.tenant_id) == 0
    assert _count(connection, 'inventory_lots', scope.tenant_id) == 0
    assert _count(connection, 'inventory_cost_layers', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200, replay.text
    assert replay.json()['import_id'] == result['import_id']
    assert replay.json()['replay'] is True
    assert _count(connection, 'inventory_items', scope.tenant_id) == 1
    assert _count(connection, 'onboarding_imports', scope.tenant_id) == 1


def test_existing_records_use_deterministic_authority_updates(client, sql_connection) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='BEANS')
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201 and first.json()['status'] == 'SUCCESS'

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['19_Inventory_Items']['C2'] = 'Updated Beans'
    workbook['19_Inventory_Items']['E2'] = '0.020000'
    workbook['21_Suppliers']['C2'] = 'Updated Supplier'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']

    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    result = updated.json()
    assert result['status'] == 'SUCCESS'
    assert result['updated'] == 2
    assert result['unchanged'] == 3
    assert _count(connection, 'inventory_items', scope.tenant_id) == 1
    assert _count(connection, 'inventory_cost_revisions', scope.tenant_id) == 2


def test_resource_import_create_update_replay_and_no_operational_session(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='TABLE', resource=True)
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['resources']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['resources']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,location_id,code,name,resource_type,status FROM resources '
            'WHERE tenant_id=%s', (scope.tenant_id,),
        )
        resource = cursor.fetchone()
    assert resource['location_id'] == scope.location_id
    assert resource['code'] == 'RESOURCE-TABLE'
    assert (resource['resource_type'], resource['status']) == ('TABLE', 'ACTIVE')
    assert _count(connection, 'restaurant_service_sessions', scope.tenant_id) == 0
    assert _count(connection, 'diner_sessions', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'resources', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['03_Resources']['C2'] = 'Updated Resource'
    workbook['03_Resources']['D2'] = 'EQUIPMENT'
    workbook['03_Resources']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['resources']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT name,resource_type,status FROM resources WHERE id=%s',
            (resource['id'],),
        )
        current = cursor.fetchone()
    assert current == {
        'name': 'Updated Resource', 'resource_type': 'EQUIPMENT',
        'status': 'INACTIVE',
    }
    assert _count(connection, 'restaurant_service_sessions', scope.tenant_id) == 0


def test_concurrent_resource_code_reconciles_and_cross_location_resolution_fails(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    with connection.cursor() as cursor:
        cursor.execute(
            'INSERT INTO locations '
            '(tenant_id,organization_id,code,name,timezone,status) '
            "VALUES (%s,%s,'OTHER','Other','America/Mexico_City','ACTIVE')",
            (scope.tenant_id, scope.organization_id),
        )
        other_location_id = cursor.lastrowid

    async def exercise() -> tuple[int, int]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        arguments = {
            'tenant_id': scope.tenant_id,
            'location_id': scope.location_id,
            'resource_code': 'CONCURRENT',
            'name': 'Concurrent Resource',
            'resource_type': 'TABLE',
            'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as first_db, database.session_factory() as second_db:
                first, second = await asyncio.gather(
                    resource_provisioning.provision_resource(first_db, **arguments),
                    resource_provisioning.provision_resource(second_db, **arguments),
                )
            async with database.session_factory() as db:
                with pytest.raises(resource_provisioning.ResourceScopeNotFoundError):
                    await resource_provisioning.resolve_resource(
                        db, tenant_id=scope.tenant_id,
                        location_id=other_location_id,
                        resource_code='CONCURRENT',
                    )
            return first.resource.id, second.resource.id
        finally:
            await database.dispose()

    first_id, second_id = asyncio.run(exercise())
    assert first_id == second_id
    assert _count(connection, 'resources', scope.tenant_id) == 1


def test_new_inactive_resource_is_plan_blocked_without_resource_write(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='INACTIVE', resource=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['03_Resources']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['status'] == 'FAILED'
    assert result['groups']['resources']['failed'] == 1
    assert result['errors'][0]['message'] == 'New Resources must be ACTIVE'
    assert _count(connection, 'resources', scope.tenant_id) == 0


def test_product_import_binding_update_replay_and_distinct_keys(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='CATALOG', product=True)
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '07_Products', (
        'PROD-SECOND', 'ORG', None, 'Product CATALOG', None, 'ACTIVE',
    ), row=3)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['products']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['products']['created'] == 2
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,name,status,source FROM products '
            'WHERE tenant_id=%s ORDER BY id',
            (scope.tenant_id,),
        )
        products = cursor.fetchall()
        cursor.execute(
            'SELECT product_id,connector_key,external_product_id '
            'FROM product_external_mappings '
            'WHERE tenant_id=%s ORDER BY external_product_id',
            (scope.tenant_id,),
        )
        bindings = cursor.fetchall()
    assert len(products) == len(bindings) == 2
    assert products[0]['id'] != products[1]['id']
    assert all(row['organization_id'] == scope.organization_id for row in products)
    assert all((row['status'], row['source']) == ('ACTIVE', 'PLATFORM') for row in products)
    expected_namespace = product_provisioning.onboarding_binding_namespace(
        contract_version=CONTRACT_VERSION, organization_id=scope.organization_id,
    )
    assert all(row['connector_key'] == expected_namespace for row in bindings)
    assert {row['external_product_id'] for row in bindings} == {'PROD-CATALOG', 'PROD-SECOND'}

    async def resolve() -> tuple[int, int]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        try:
            async with database.session_factory() as db:
                namespace = product_provisioning.onboarding_binding_namespace(
                    contract_version=CONTRACT_VERSION,
                    organization_id=scope.organization_id,
                )
                first_product = await product_provisioning.resolve_product_binding(
                    db, tenant_id=scope.tenant_id, organization_id=scope.organization_id,
                    binding_namespace=namespace, product_key='PROD-CATALOG',
                )
                second_product = await product_provisioning.resolve_product_binding(
                    db, tenant_id=scope.tenant_id, organization_id=scope.organization_id,
                    binding_namespace=namespace, product_key='PROD-SECOND',
                )
                return first_product.id, second_product.id
        finally:
            await database.dispose()

    resolved = asyncio.run(resolve())
    assert set(resolved) == {row['id'] for row in products}

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'products', scope.tenant_id) == 2

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['07_Products']['D2'] = 'Updated canonical Product'
    changed_output = BytesIO()
    workbook.save(changed_output)
    workbook.close()
    changed = changed_output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['products']['updated'] == 1
    assert updated.json()['groups']['products']['unchanged'] == 1
    assert _count(connection, 'products', scope.tenant_id) == 2


def test_category_import_precedes_and_resolves_explicit_product_category(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='MENU', category_product=True,
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '06_Categories', (
        'CAT-CHILD', 'ORG', 'Child', 'CAT-MENU', 2, 'ACTIVE',
    ), row=2)
    _set_row(workbook, '06_Categories', (
        'CAT-MENU', 'ORG', 'Menu', None, 1, 'ACTIVE',
    ), row=3)
    workbook['07_Products']['C2'] = 'CAT-CHILD'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()

    analysis = _analyze(client, headers, content)
    assert analysis.status_code == 200
    assert analysis.json()['analysis']['status'] == 'VALID'
    fingerprint = analysis.json()['dataset_fingerprint']
    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['categories']['created'] == 2
    assert result['groups']['products']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,parent_id,name FROM product_categories '
            'WHERE tenant_id=%s ORDER BY name', (scope.tenant_id,),
        )
        categories = cursor.fetchall()
        cursor.execute(
            'SELECT category_id FROM products WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        product = cursor.fetchone()
    by_name = {row['name']: row for row in categories}
    assert by_name['Child']['parent_id'] == by_name['Menu']['id']
    assert product['category_id'] == by_name['Child']['id']
    assert _count(connection, 'product_category_external_mappings', scope.tenant_id) == 2

    async def resolve_category() -> int:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        try:
            async with database.session_factory() as db:
                category = await category_provisioning.resolve_category_binding(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=category_provisioning.onboarding_binding_namespace(
                        contract_version=CONTRACT_VERSION,
                        organization_id=scope.organization_id,
                    ),
                    category_key='CAT-CHILD',
                )
                return category.id
        finally:
            await database.dispose()

    assert asyncio.run(resolve_category()) == by_name['Child']['id']
    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'product_categories', scope.tenant_id) == 2

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['06_Categories']['C2'] = 'Renamed Child'
    workbook['06_Categories']['E2'] = 4
    changed_output = BytesIO()
    workbook.save(changed_output)
    workbook.close()
    changed = changed_output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['categories']['updated'] == 1
    assert updated.json()['groups']['categories']['unchanged'] == 1
    assert updated.json()['groups']['products']['unchanged'] == 1
    assert _count(connection, 'product_categories', scope.tenant_id) == 2


def test_category_name_conflict_fails_without_ambiguous_binding(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='CATEGORY')
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '06_Categories', (
        'CAT-A', 'ORG', 'Same category', None, 0, 'ACTIVE',
    ))
    _set_row(workbook, '06_Categories', (
        'CAT-B', 'ORG', 'Same category', None, 0, 'ACTIVE',
    ), row=3)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['groups']['categories']['created'] == 1
    assert result['groups']['categories']['failed'] == 1
    assert _count(connection, 'product_categories', scope.tenant_id) == 1
    assert _count(connection, 'product_category_external_mappings', scope.tenant_id) == 1


def test_concurrent_category_binding_reconciles_and_rejects_cross_scope_resolution(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)

    async def exercise() -> tuple[int, int]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = category_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'binding_namespace': namespace,
            'category_key': 'CONCURRENT',
            'parent_category_key': None,
            'name': 'Concurrent category',
            'display_order': 0,
            'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as first_db, database.session_factory() as second_db:
                first, second = await asyncio.gather(
                    category_provisioning.provision_category(first_db, **arguments),
                    category_provisioning.provision_category(second_db, **arguments),
                )
            async with database.session_factory() as db:
                with pytest.raises(category_provisioning.CategoryScopeNotFoundError):
                    await category_provisioning.resolve_category_binding(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id + 999999,
                        binding_namespace=namespace, category_key='CONCURRENT',
                    )
            return first.category.id, second.category.id
        finally:
            await database.dispose()

    first_id, second_id = asyncio.run(exercise())
    assert first_id == second_id
    assert _count(connection, 'product_categories', scope.tenant_id) == 1
    assert _count(connection, 'product_category_external_mappings', scope.tenant_id) == 1


def test_product_binding_conflict_across_organization_blocks_before_product_write(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    with connection.cursor() as cursor:
        cursor.execute(
            "INSERT INTO organizations (tenant_id,code,name,status) VALUES (%s,'OTHER','Other','ACTIVE')",
            (scope.tenant_id,),
        )
        other_organization_id = cursor.lastrowid
        cursor.execute(
            "INSERT INTO products (tenant_id,organization_id,name,status,source) "
            "VALUES (%s,%s,'Foreign','ACTIVE','PLATFORM')",
            (scope.tenant_id, other_organization_id),
        )
        foreign_product_id = cursor.lastrowid
        namespace = product_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION, organization_id=scope.organization_id,
        )
        cursor.execute(
            'INSERT INTO product_external_mappings '
            '(tenant_id,product_id,connector_key,external_product_id) VALUES (%s,%s,%s,%s)',
            (scope.tenant_id, foreign_product_id, namespace, 'PROD-CONFLICT'),
        )

    content = _workbook(f'{prefix}-onboarding', item_code='CONFLICT', product=True)
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['status'] == 'FAILED'
    assert result['groups']['products']['failed'] == 1
    assert result['errors'][0]['code'] == 'PLAN_BLOCKED'
    assert _count(connection, 'products', scope.tenant_id) == 1


def test_partial_authority_failure_is_durable_and_not_blindly_replayed(
    client, sql_connection, monkeypatch,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='FAILSAFE')
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    async def reject_conversion(*args, **kwargs):
        raise RuntimeError('simulated authority failure')

    monkeypatch.setattr(inventory_service, 'append_item_uom_conversion', reject_conversion)
    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['status'] == 'PARTIAL'
    assert result['created'] == 1 and result['failed'] == 1
    assert result['groups']['suppliers']['created'] == 0
    assert result['errors'][0]['message'] == 'An existing write authority rejected the planned operation'

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200
    assert replay.json()['import_id'] == result['import_id']
    assert replay.json()['status'] == 'PARTIAL' and replay.json()['replay'] is True
    assert _count(connection, 'inventory_items', scope.tenant_id) == 1
    assert _count(connection, 'onboarding_imports', scope.tenant_id) == 1


def test_invalid_and_cross_scope_confirmation_cannot_persist(client, sql_connection) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    invalid = bytearray(_workbook(f'{prefix}-onboarding', item_code='INVALID'))
    invalid[-8:] = b'not-xlsx'
    invalid_analysis = _analyze(client, headers, bytes(invalid)).json()
    response = _confirm(
        client, headers, bytes(invalid), scope.location_id,
        invalid_analysis['dataset_fingerprint'],
    )
    assert response.status_code == 422

    content = _workbook(f'{prefix}-onboarding', item_code='CROSS')
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    response = _confirm(client, headers, content, scope.location_id + 999999, fingerprint)
    assert response.status_code == 403
    assert _count(connection, 'inventory_items', scope.tenant_id) == 0
    assert _count(connection, 'onboarding_imports', scope.tenant_id) == 0


def test_preparation_configuration_import_create_update_replay_without_areas(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='RICE',
        preparation_configuration=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['groups']['preparation_configuration']['created'] == 1
    assert result['groups']['preparation_configuration'] == {
        'classification': 'IMPORTABLE_NOW',
        'authority': (
            'restaurant.preparation.configuration_provisioning.'
            'provision_configuration'
        ),
        'required_for_stage0': True,
        'planned': 1, 'created': 1, 'updated': 0, 'unchanged': 0,
        'deferred': 0, 'failed': 0,
    }
    assert _count(connection, 'location_preparation_configurations', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 0
    assert _count(connection, 'preparation_works', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200
    assert replay.json()['replay'] is True
    assert replay.json()['import_id'] == result['import_id']
    assert _count(connection, 'location_preparation_configurations', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['04_Prep_Config']['B2'] = 'EXTERNAL_POS'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(
        client, headers, changed, scope.location_id, changed_fingerprint,
    )
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['preparation_configuration']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT organization_id,location_id,preparation_owner '
            'FROM location_preparation_configurations WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        configuration = cursor.fetchone()
    assert configuration == {
        'organization_id': scope.organization_id,
        'location_id': scope.location_id,
        'preparation_owner': 'EXTERNAL_POS',
    }
    assert _count(connection, 'location_preparation_configurations', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 0


def test_preparation_configuration_authority_is_scoped_and_concurrency_safe(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'preparation_owner': 'PLATFORM',
        }
        try:
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    configuration_provisioning.provision_configuration(
                        first_db, **arguments,
                    ),
                    configuration_provisioning.provision_configuration(
                        second_db, **arguments,
                    ),
                )
            async with database.session_factory() as db:
                unchanged = await configuration_provisioning.plan_configuration(
                    db, **arguments,
                )
                assert unchanged.operation == 'UNCHANGED'
                with pytest.raises(
                    configuration_provisioning.PreparationConfigurationScopeNotFoundError
                ):
                    await configuration_provisioning.resolve_configuration(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id + 999999,
                        location_id=scope.location_id,
                    )
                with pytest.raises(
                    configuration_provisioning.PreparationConfigurationProvisioningError
                ):
                    await configuration_provisioning.plan_configuration(
                        db, **{**arguments, 'preparation_owner': 'INVALID'},
                    )
            return (
                first.configuration.id, second.configuration.id,
                {first.operation, second.operation},
            )
        finally:
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'location_preparation_configurations', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 0


def test_preparation_area_import_resolves_resource_updates_and_replays(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='BAR', resource=True,
        preparation_configuration=True, preparation_area=True,
        area_resource=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['groups']['preparation_areas'] == {
        'classification': 'IMPORTABLE_NOW',
        'authority': 'restaurant.preparation.area_provisioning.provision_area',
        'required_for_stage0': True,
        'planned': 1, 'created': 1, 'updated': 0, 'unchanged': 0,
        'deferred': 0, 'failed': 0,
    }
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT a.id,a.organization_id,a.location_id,a.code,a.name,a.status,'
            'r.code AS resource_code FROM preparation_areas a '
            'LEFT JOIN resources r ON r.id=a.resource_id AND r.tenant_id=a.tenant_id '
            'WHERE a.tenant_id=%s', (scope.tenant_id,),
        )
        area = cursor.fetchone()
    assert area == {
        'id': area['id'], 'organization_id': scope.organization_id,
        'location_id': scope.location_id, 'code': 'AREA-BAR',
        'name': 'Area BAR', 'status': 'ACTIVE',
        'resource_code': 'RESOURCE-BAR',
    }
    assert _count(connection, 'resources', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 1
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 0
    assert _count(connection, 'preparation_works', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200
    assert replay.json()['replay'] is True
    assert replay.json()['import_id'] == result['import_id']
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['05_Prep_Areas']['C2'] = 'Updated Bar'
    workbook['05_Prep_Areas']['D2'] = None
    workbook['05_Prep_Areas']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(
        client, headers, changed, scope.location_id, changed_fingerprint,
    )
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['preparation_areas']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT name,status,resource_id FROM preparation_areas WHERE id=%s',
            (area['id'],),
        )
        current = cursor.fetchone()
    assert current == {
        'name': 'Updated Bar', 'status': 'INACTIVE', 'resource_id': None,
    }
    assert _count(connection, 'resources', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 1
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 0


def test_preparation_area_authority_is_scoped_convergent_and_never_creates_resource(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    with connection.cursor() as cursor:
        cursor.execute(
            'INSERT INTO locations '
            '(tenant_id,organization_id,code,name,timezone,status) '
            "VALUES (%s,%s,'OTHER','Other','America/Mexico_City','ACTIVE')",
            (scope.tenant_id, scope.organization_id),
        )
        other_location_id = cursor.lastrowid
        cursor.execute(
            'INSERT INTO resources '
            '(tenant_id,location_id,code,name,resource_type,status) '
            "VALUES (%s,%s,'FOREIGN','Foreign','AREA','ACTIVE')",
            (scope.tenant_id, other_location_id),
        )

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'area_code': 'BAR',
            'resource_code': None,
            'name': 'Bar',
            'status': 'ACTIVE',
        }
        try:
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    area_provisioning.provision_area(first_db, **arguments),
                    area_provisioning.provision_area(second_db, **arguments),
                )
            async with database.session_factory() as db:
                resolved = await area_provisioning.resolve_area(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    location_id=scope.location_id, area_code='BAR',
                )
                unchanged = await area_provisioning.plan_area(db, **arguments)
                assert resolved.id == first.area.id
                assert unchanged.operation == 'UNCHANGED'
                with pytest.raises(
                    area_provisioning.PreparationAreaScopeNotFoundError
                ):
                    await area_provisioning.resolve_area(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id,
                        location_id=other_location_id, area_code='BAR',
                    )
                with pytest.raises(
                    area_provisioning.PreparationAreaScopeNotFoundError
                ):
                    await area_provisioning.plan_area(
                        db, **{**arguments, 'area_code': 'FOREIGN-RESOURCE',
                               'resource_code': 'FOREIGN'},
                    )
                with pytest.raises(area_provisioning.PreparationAreaConflictError):
                    await area_provisioning.plan_area(
                        db, **{**arguments, 'area_code': 'NEW-INACTIVE',
                               'status': 'INACTIVE'},
                    )
            return (
                first.area.id, second.area.id,
                {first.operation, second.operation},
            )
        finally:
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'resources', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 1
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 0


def test_preparation_route_import_preserves_history_and_replays(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='ROUTE', category_product=True,
        preparation_configuration=True, preparation_area=True,
        preparation_route=True,
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '05_Prep_Areas', (
        'LOC', 'AREA-SECOND', 'Second Area', None, 'ACTIVE',
    ), row=3)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    response = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result['groups']['preparation_routes'] == {
        'classification': 'IMPORTABLE_NOW',
        'authority': (
            'restaurant.preparation.route_provisioning.provision_route_binding'
        ),
        'required_for_stage0': True,
        'planned': 1, 'created': 1, 'updated': 0, 'unchanged': 0,
        'deferred': 0, 'failed': 0,
    }
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT r.id,r.policy,r.status,r.active_slot,a.code AS area_code '
            'FROM product_preparation_routes r '
            'LEFT JOIN preparation_areas a ON a.id=r.preparation_area_id '
            'WHERE r.tenant_id=%s ORDER BY r.id', (scope.tenant_id,),
        )
        first_routes = cursor.fetchall()
    assert len(first_routes) == 1
    assert first_routes[0] == {
        'id': first_routes[0]['id'], 'policy': 'AREA', 'status': 'ACTIVE',
        'active_slot': 1, 'area_code': 'AREA-ROUTE',
    }
    original_route_id = first_routes[0]['id']
    assert _count(connection, 'preparation_works', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200
    assert replay.json()['replay'] is True
    assert replay.json()['import_id'] == result['import_id']
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['05_Prep_Areas']['C2'] = 'Renamed Route Area'
    unchanged_output = BytesIO()
    workbook.save(unchanged_output)
    workbook.close()
    unchanged_content = unchanged_output.getvalue()
    unchanged_fingerprint = _analyze(
        client, headers, unchanged_content,
    ).json()['dataset_fingerprint']
    unchanged = _confirm(
        client, headers, unchanged_content, scope.location_id,
        unchanged_fingerprint,
    )
    assert unchanged.status_code == 201, unchanged.text
    assert unchanged.json()['groups']['preparation_routes']['unchanged'] == 1
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(unchanged_content), data_only=False)
    workbook['27_Prep_Routes']['D2'] = 'AREA-SECOND'
    changed_output = BytesIO()
    workbook.save(changed_output)
    workbook.close()
    changed = changed_output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(
        client, headers, changed, scope.location_id, changed_fingerprint,
    )
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['preparation_routes']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT r.id,r.policy,r.status,r.active_slot,a.code AS area_code '
            'FROM product_preparation_routes r '
            'LEFT JOIN preparation_areas a ON a.id=r.preparation_area_id '
            'WHERE r.tenant_id=%s ORDER BY r.id', (scope.tenant_id,),
        )
        routes = cursor.fetchall()
    assert len(routes) == 2
    assert routes[0] == {
        'id': original_route_id, 'policy': 'AREA', 'status': 'INACTIVE',
        'active_slot': None, 'area_code': 'AREA-ROUTE',
    }
    assert routes[1]['policy'] == 'AREA'
    assert routes[1]['status'] == 'ACTIVE'
    assert routes[1]['active_slot'] == 1
    assert routes[1]['area_code'] == 'AREA-SECOND'
    assert _count(connection, 'preparation_works', scope.tenant_id) == 0


def test_preparation_route_bindings_are_scoped_and_concurrency_safe(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='CONCURRENT',
        category_product=True, preparation_area=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    prepared = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert prepared.status_code == 201, prepared.text
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 0

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = product_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'binding_namespace': namespace,
            'product_key': 'PROD-CONCURRENT',
            'policy': 'AREA',
            'preparation_area_code': 'AREA-CONCURRENT',
            'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as db:
                with pytest.raises(
                    route_provisioning.PreparationRouteScopeNotFoundError
                ):
                    await route_provisioning.plan_route_binding(
                        db, **{**arguments, 'product_key': 'UNKNOWN'},
                    )
                with pytest.raises(
                    route_provisioning.PreparationRouteScopeNotFoundError
                ):
                    await route_provisioning.plan_route_binding(
                        db, **{**arguments,
                               'location_id': scope.location_id + 999999},
                    )
                with pytest.raises(
                    route_provisioning.PreparationRouteScopeNotFoundError
                ):
                    await route_provisioning.plan_route_binding(
                        db, **{**arguments,
                               'preparation_area_code': 'UNKNOWN'},
                    )
                with pytest.raises(
                    route_provisioning.PreparationRouteConflictError
                ):
                    await route_provisioning.plan_route_binding(
                        db, **{**arguments, 'status': 'INACTIVE'},
                    )
                with pytest.raises(
                    route_provisioning.PreparationRouteProvisioningError
                ):
                    await route_provisioning.plan_route_binding(
                        db, **{**arguments, 'policy': 'COMPONENTS'},
                    )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    route_provisioning.provision_route_binding(
                        first_db, **arguments,
                    ),
                    route_provisioning.provision_route_binding(
                        second_db, **arguments,
                    ),
                )
            return (
                first.route.id, second.route.id,
                {first.operation, second.operation},
            )
        finally:
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'products', scope.tenant_id) == 1
    assert _count(connection, 'preparation_areas', scope.tenant_id) == 1
    assert _count(connection, 'product_preparation_routes', scope.tenant_id) == 1
    assert _count(connection, 'preparation_works', scope.tenant_id) == 0


def test_menu_import_binding_update_replay_without_dependent_catalog_writes(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(f'{prefix}-onboarding', item_code='DINNER', menu=True)
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['menus']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['menus']['authority'] == (
        'restaurant.menu_provisioning.provision_menu'
    )
    assert result['groups']['menus']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,name,status FROM menus WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        menu = cursor.fetchone()
        cursor.execute(
            'SELECT menu_id,connector_key,external_menu_id '
            'FROM menu_external_mappings WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        binding = cursor.fetchone()
        cursor.execute(
            'SELECT menu_id,location_id,status FROM menu_locations '
            'WHERE tenant_id=%s', (scope.tenant_id,),
        )
        assignment = cursor.fetchone()
    assert menu['organization_id'] == scope.organization_id
    assert (menu['name'], menu['status']) == ('Menu DINNER', 'ACTIVE')
    assert binding['menu_id'] == assignment['menu_id'] == menu['id']
    assert binding['external_menu_id'] == 'MENU-DINNER'
    assert assignment == {
        'menu_id': menu['id'], 'location_id': scope.location_id,
        'status': 'ACTIVE',
    }
    assert binding['connector_key'] == menu_provisioning.onboarding_binding_namespace(
        contract_version=CONTRACT_VERSION,
        organization_id=scope.organization_id,
    )
    assert _count(connection, 'menu_sections', scope.tenant_id) == 0
    assert _count(connection, 'menu_items', scope.tenant_id) == 0
    assert _count(connection, 'product_prices', scope.tenant_id) == 0

    async def resolve() -> int:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        try:
            async with database.session_factory() as db:
                resolved = await menu_provisioning.resolve_menu_binding(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=binding['connector_key'],
                    menu_key='MENU-DINNER',
                )
                return resolved.id
        finally:
            await database.dispose()

    assert asyncio.run(resolve()) == menu['id']
    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'menus', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['08_Menus']['C2'] = 'Dinner updated'
    workbook['08_Menus']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['menus']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute('SELECT name,status FROM menus WHERE id=%s', (menu['id'],))
        current_menu = cursor.fetchone()
        cursor.execute(
            'SELECT status FROM menu_locations WHERE menu_id=%s AND location_id=%s',
            (menu['id'], scope.location_id),
        )
        current_assignment = cursor.fetchone()
    assert current_menu == {'name': 'Dinner updated', 'status': 'ACTIVE'}
    assert current_assignment == {'status': 'INACTIVE'}
    assert _count(connection, 'menus', scope.tenant_id) == 1
    assert _count(connection, 'menu_external_mappings', scope.tenant_id) == 1
    assert _count(connection, 'menu_sections', scope.tenant_id) == 0
    assert _count(connection, 'menu_items', scope.tenant_id) == 0
    assert _count(connection, 'product_prices', scope.tenant_id) == 0


def test_menu_authority_is_scoped_concurrency_safe_and_rejects_unsafe_creation(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = menu_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'binding_namespace': namespace,
            'menu_key': 'CONCURRENT',
            'name': 'Concurrent Menu',
            'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as db:
                with pytest.raises(menu_provisioning.MenuConflictError):
                    await menu_provisioning.plan_menu(
                        db, **{**arguments, 'menu_key': 'INACTIVE',
                               'status': 'INACTIVE'},
                    )
                with pytest.raises(menu_provisioning.MenuScopeNotFoundError):
                    await menu_provisioning.plan_menu(
                        db, **{**arguments,
                               'location_id': scope.location_id + 999999},
                    )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    menu_provisioning.provision_menu(first_db, **arguments),
                    menu_provisioning.provision_menu(second_db, **arguments),
                )
            async with database.session_factory() as db:
                with pytest.raises(menu_provisioning.MenuScopeNotFoundError):
                    await menu_provisioning.resolve_menu_binding(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id + 999999,
                        binding_namespace=namespace, menu_key='CONCURRENT',
                    )
            return first.menu.id, second.menu.id, {
                first.operation, second.operation,
            }
        finally:
            await database.dispose()

    with pytest.raises(menu_provisioning.MenuConflictError):
        menu_provisioning.validate_menu_definitions((
            ('DUPLICATE', 'First name'), ('DUPLICATE', 'Second name'),
        ))
    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'menus', scope.tenant_id) == 1
    assert _count(connection, 'menu_external_mappings', scope.tenant_id) == 1
    assert _count(connection, 'menu_locations', scope.tenant_id) == 1
    assert _count(connection, 'menu_sections', scope.tenant_id) == 0
    assert _count(connection, 'menu_items', scope.tenant_id) == 0


def test_menu_section_import_orders_dependencies_updates_and_resolves_item_inputs(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='SECTIONS', product=True,
        menu_section=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['menus']['created'] == 1
    assert result['groups']['menu_sections']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['menu_sections']['authority'] == (
        'restaurant.menu_section_provisioning.provision_menu_section'
    )
    assert result['groups']['menu_sections']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,menu_id,name,display_order,status '
            'FROM menu_sections WHERE tenant_id=%s', (scope.tenant_id,),
        )
        section = cursor.fetchone()
        cursor.execute(
            'SELECT menu_id,section_id,connector_key,external_section_id '
            'FROM menu_section_external_mappings WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        binding = cursor.fetchone()
    assert section['organization_id'] == scope.organization_id
    assert (section['name'], section['display_order'], section['status']) == (
        'Section SECTIONS', 10, 'ACTIVE',
    )
    assert binding['menu_id'] == section['menu_id']
    assert binding['section_id'] == section['id']
    assert binding['external_section_id'] == 'SECTION-SECTIONS'
    assert _count(connection, 'menu_items', scope.tenant_id) == 0
    assert _count(connection, 'product_prices', scope.tenant_id) == 0

    async def resolve_item_inputs() -> tuple[int, int, int]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = menu_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        try:
            async with database.session_factory() as db:
                menu = await menu_provisioning.resolve_menu_binding(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=namespace, menu_key='MENU-SECTIONS',
                )
                resolved_section = (
                    await menu_section_provisioning.resolve_menu_section_binding(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id,
                        binding_namespace=namespace, menu_key='MENU-SECTIONS',
                        section_key='SECTION-SECTIONS',
                    )
                )
                product = await product_provisioning.resolve_product_binding(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=namespace, product_key='PROD-SECTIONS',
                )
                return menu.id, resolved_section.id, product.id
        finally:
            await database.dispose()

    menu_id, section_id, product_id = asyncio.run(resolve_item_inputs())
    assert menu_id == section['menu_id']
    assert section_id == section['id']
    assert product_id > 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'menu_sections', scope.tenant_id) == 1
    assert _count(connection, 'menu_section_external_mappings', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['09_Menu_Sections']['C2'] = 'Updated Section'
    workbook['09_Menu_Sections']['D2'] = 30
    workbook['09_Menu_Sections']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['menus']['unchanged'] == 1
    assert updated.json()['groups']['menu_sections']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT name,display_order,status FROM menu_sections WHERE id=%s',
            (section['id'],),
        )
        current = cursor.fetchone()
    assert current == {
        'name': 'Updated Section', 'display_order': 30, 'status': 'INACTIVE',
    }
    assert _count(connection, 'menu_sections', scope.tenant_id) == 1
    assert _count(connection, 'menu_items', scope.tenant_id) == 0
    assert _count(connection, 'product_prices', scope.tenant_id) == 0


def test_menu_section_binding_is_menu_scoped_and_concurrency_safe(
    sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)

    async def exercise() -> tuple[int, int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = menu_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        menu_arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'binding_namespace': namespace,
            'name': 'Menu', 'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as db:
                await menu_provisioning.provision_menu(
                    db, **menu_arguments, menu_key='FIRST',
                )
                await menu_provisioning.provision_menu(
                    db, **{**menu_arguments, 'menu_key': 'SECOND',
                           'name': 'Second Menu'},
                )
            section_arguments = {
                'tenant_id': scope.tenant_id,
                'organization_id': scope.organization_id,
                'binding_namespace': namespace,
                'menu_key': 'FIRST', 'section_key': 'SHARED',
                'name': 'Shared Section', 'display_order': 1,
                'status': 'ACTIVE',
            }
            async with database.session_factory() as db:
                with pytest.raises(menu_section_provisioning.MenuSectionConflictError):
                    await menu_section_provisioning.plan_menu_section(
                        db, **{**section_arguments, 'section_key': 'INACTIVE',
                               'status': 'INACTIVE'},
                    )
                with pytest.raises(
                    menu_section_provisioning.MenuSectionScopeNotFoundError
                ):
                    await menu_section_provisioning.plan_menu_section(
                        db, **{**section_arguments, 'menu_key': 'UNKNOWN'},
                    )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    menu_section_provisioning.provision_menu_section(
                        first_db, **section_arguments,
                    ),
                    menu_section_provisioning.provision_menu_section(
                        second_db, **section_arguments,
                    ),
                )
            async with database.session_factory() as db:
                other_menu_section = (
                    await menu_section_provisioning.provision_menu_section(
                        db, **{**section_arguments, 'menu_key': 'SECOND'},
                    )
                )
                with pytest.raises(
                    menu_section_provisioning.MenuSectionScopeNotFoundError
                ):
                    await menu_section_provisioning.resolve_menu_section_binding(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id + 999999,
                        binding_namespace=namespace, menu_key='FIRST',
                        section_key='SHARED',
                    )
            return (
                first.section.id, second.section.id,
                other_menu_section.section.id,
                {first.operation, second.operation},
            )
        finally:
            await database.dispose()

    first_id, second_id, other_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert other_id != first_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'menus', scope.tenant_id) == 2
    assert _count(connection, 'menu_sections', scope.tenant_id) == 2
    assert _count(connection, 'menu_section_external_mappings', scope.tenant_id) == 2
    assert _count(connection, 'menu_items', scope.tenant_id) == 0
    assert _count(connection, 'product_prices', scope.tenant_id) == 0


def test_menu_item_import_uses_existing_identity_updates_and_never_prices(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='PLACEMENT', menu_item=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert result['groups']['products']['created'] == 1
    assert result['groups']['menus']['created'] == 1
    assert result['groups']['menu_sections']['created'] == 1
    assert result['groups']['menu_items']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['menu_items']['authority'] == (
        'restaurant.menu_item_provisioning.provision_menu_item'
    )
    assert result['groups']['menu_items']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,menu_id,section_id,product_id,'
            'display_order,status FROM menu_items WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        item = cursor.fetchone()
    assert item['organization_id'] == scope.organization_id
    assert (item['display_order'], item['status']) == (20, 'ACTIVE')
    assert _count(connection, 'menu_items', scope.tenant_id) == 1
    assert _count(connection, 'product_prices', scope.tenant_id) == 0

    async def resolve() -> int:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = menu_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        try:
            async with database.session_factory() as db:
                resolved = await menu_item_provisioning.resolve_menu_item(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=namespace, menu_key='MENU-PLACEMENT',
                    section_key='SECTION-PLACEMENT',
                    product_key='PROD-PLACEMENT',
                )
                return resolved.id
        finally:
            await database.dispose()

    assert asyncio.run(resolve()) == item['id']
    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'menu_items', scope.tenant_id) == 1
    assert _count(connection, 'product_prices', scope.tenant_id) == 0

    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '09_Menu_Sections', (
        'SECTION-SECOND', 'MENU-PLACEMENT', 'Second Section', 30, 'ACTIVE',
    ), row=3)
    workbook['10_Menu_Items']['B2'] = 'SECTION-SECOND'
    workbook['10_Menu_Items']['D2'] = 40
    workbook['10_Menu_Items']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['menu_sections']['created'] == 1
    assert updated.json()['groups']['menu_items']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT section_id,display_order,status FROM menu_items WHERE id=%s',
            (item['id'],),
        )
        current = cursor.fetchone()
        cursor.execute(
            'SELECT id FROM menu_sections WHERE tenant_id=%s AND name=%s',
            (scope.tenant_id, 'Second Section'),
        )
        second_section = cursor.fetchone()
    assert current == {
        'section_id': second_section['id'], 'display_order': 40,
        'status': 'INACTIVE',
    }
    assert _count(connection, 'menu_items', scope.tenant_id) == 1
    assert _count(connection, 'product_prices', scope.tenant_id) == 0


def test_menu_item_authority_rejects_foreign_section_and_converges_concurrently(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='CONCURRENT-ITEM',
        product=True, menu_section=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    seeded = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert seeded.status_code == 201, seeded.text
    assert _count(connection, 'menu_items', scope.tenant_id) == 0

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = menu_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'binding_namespace': namespace,
            'menu_key': 'MENU-CONCURRENT-ITEM',
            'section_key': 'SECTION-CONCURRENT-ITEM',
            'product_key': 'PROD-CONCURRENT-ITEM',
            'display_order': 5, 'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as db:
                await menu_provisioning.provision_menu(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    location_id=scope.location_id,
                    binding_namespace=namespace, menu_key='OTHER-MENU',
                    name='Other Menu', status='ACTIVE',
                )
                await menu_section_provisioning.provision_menu_section(
                    db, tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    binding_namespace=namespace, menu_key='OTHER-MENU',
                    section_key='FOREIGN-SECTION', name='Foreign Section',
                    display_order=0, status='ACTIVE',
                )
                with pytest.raises(
                    menu_item_provisioning.MenuItemScopeNotFoundError
                ):
                    await menu_item_provisioning.plan_menu_item(
                        db, **{**arguments, 'section_key': 'FOREIGN-SECTION'},
                    )
                with pytest.raises(
                    menu_item_provisioning.MenuItemScopeNotFoundError
                ):
                    await menu_item_provisioning.plan_menu_item(
                        db, **{**arguments,
                               'organization_id': scope.organization_id + 999999},
                    )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    menu_item_provisioning.provision_menu_item(
                        first_db, **arguments,
                    ),
                    menu_item_provisioning.provision_menu_item(
                        second_db, **arguments,
                    ),
                )
            return first.menu_item.id, second.menu_item.id, {
                first.operation, second.operation,
            }
        finally:
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'menu_items', scope.tenant_id) == 1
    assert _count(connection, 'product_prices', scope.tenant_id) == 0


def test_price_import_preserves_exact_money_updates_and_avoids_fiscal_writes(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='PRICE', price=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert not result['errors'], result['errors']
    assert result['groups']['products']['created'] == 1, result
    assert result['groups']['prices']['classification'] == 'IMPORTABLE_NOW'
    assert result['groups']['prices']['authority'] == (
        'restaurant.pricing.provisioning.provision_price'
    )
    assert result['groups']['prices']['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,product_id,location_id,amount,currency,'
            'status,source FROM product_prices WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        price = cursor.fetchone()
    assert price['organization_id'] == scope.organization_id
    assert price['location_id'] == scope.location_id
    assert price['amount'] == Decimal('12.3400')
    assert (price['currency'], price['status'], price['source']) == (
        'MXN', 'ACTIVE', 'PLATFORM',
    )
    assert _count(connection, 'product_prices', scope.tenant_id) == 1
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 0
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 0

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'product_prices', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['11_Prices']['C2'] = '99.8765'
    workbook['11_Prices']['D2'] = 'USD'
    workbook['11_Prices']['E2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['prices']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,amount,currency,status,source FROM product_prices WHERE id=%s',
            (price['id'],),
        )
        current = cursor.fetchone()
    assert current == {
        'id': price['id'], 'amount': Decimal('99.8765'), 'currency': 'USD',
        'status': 'INACTIVE', 'source': 'PLATFORM',
    }
    assert _count(connection, 'product_prices', scope.tenant_id) == 1
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 0
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 0


def test_price_authority_validates_scope_source_and_converges_concurrently(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='PRICE-RACE', product=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    seeded = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert seeded.status_code == 201, seeded.text
    assert _count(connection, 'product_prices', scope.tenant_id) == 0

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = product_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'binding_namespace': namespace,
            'product_key': 'PROD-PRICE-RACE',
            'amount': Decimal('7.1234'), 'currency': 'MXN',
            'status': 'ACTIVE',
        }
        try:
            async with database.session_factory() as db:
                with pytest.raises(price_provisioning.PriceScopeNotFoundError):
                    await price_provisioning.plan_price(
                        db, **{**arguments,
                               'location_id': scope.location_id + 999999},
                    )
                for changes in (
                    {'amount': Decimal('1.00001')},
                    {'currency': 'MÉX'},
                    {'status': 'BROKEN'},
                ):
                    with pytest.raises(price_provisioning.PriceProvisioningError):
                        await price_provisioning.plan_price(
                            db, **{**arguments, **changes},
                        )
                with pytest.raises(price_provisioning.PriceConflictError):
                    await price_provisioning.plan_price(
                        db, **{**arguments, 'status': 'INACTIVE'},
                    )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    price_provisioning.provision_price(first_db, **arguments),
                    price_provisioning.provision_price(second_db, **arguments),
                )
            return first.price.id, second.price.id, {
                first.operation, second.operation,
            }
        finally:
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'product_prices', scope.tenant_id) == 1
    with connection.cursor() as cursor:
        cursor.execute(
            "UPDATE product_prices SET source='POS' WHERE id=%s", (first_id,),
        )

    async def reject_pos() -> None:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        try:
            async with database.session_factory() as db:
                with pytest.raises(price_provisioning.PriceConflictError):
                    await price_provisioning.plan_price(
                        db, tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id,
                        location_id=scope.location_id,
                        binding_namespace=(
                            product_provisioning.onboarding_binding_namespace(
                                contract_version=CONTRACT_VERSION,
                                organization_id=scope.organization_id,
                            )
                        ),
                        product_key='PROD-PRICE-RACE',
                        amount=Decimal('7.1234'), currency='MXN',
                        status='ACTIVE',
                    )
        finally:
            await database.dispose()

    asyncio.run(reject_pos())
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 0
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 0


def test_warehouse_import_updates_replays_and_creates_no_inventory_evidence(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='WAREHOUSE',
        inventory=False, warehouse=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert not result['errors'], result['errors']
    assert result['groups']['warehouses'] == {
        'classification': 'IMPORTABLE_NOW',
        'authority': (
            'restaurant.inventory.warehouse_provisioning.provision_warehouse'
        ),
        'required_for_stage0': True,
        'planned': 1, 'created': 1, 'updated': 0, 'unchanged': 0,
        'deferred': 0, 'failed': 0,
    }
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,organization_id,location_id,code,name,status,default_slot,'
            'negative_stock_policy,version FROM warehouses WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        warehouse = cursor.fetchone()
    assert warehouse == {
        'id': warehouse['id'], 'organization_id': scope.organization_id,
        'location_id': scope.location_id, 'code': 'WH-WAREHOUSE',
        'name': 'Warehouse WAREHOUSE', 'status': 'ACTIVE',
        'default_slot': 1, 'negative_stock_policy': 'ALLOW', 'version': 1,
    }

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'warehouses', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['18_Warehouses']['C2'] = 'Updated Warehouse'
    workbook['18_Warehouses']['E2'] = 'BLOCK'
    workbook['18_Warehouses']['F2'] = 'INACTIVE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    assert updated.json()['groups']['warehouses']['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT id,name,status,default_slot,negative_stock_policy,version '
            'FROM warehouses WHERE tenant_id=%s', (scope.tenant_id,),
        )
        current = cursor.fetchone()
    assert current == {
        'id': warehouse['id'], 'name': 'Updated Warehouse',
        'status': 'INACTIVE', 'default_slot': 1,
        'negative_stock_policy': 'BLOCK', 'version': 2,
    }
    for table in (
        'inventory_items', 'stock_movements', 'inventory_lots',
        'inventory_cost_layers', 'inventory_transfers', 'goods_receipts',
        'inventory_losses', 'physical_counts', 'preparation_batches',
        'inventory_valuation_snapshots',
    ):
        assert _count(connection, table, scope.tenant_id) == 0, table


def test_warehouse_authority_rejects_scope_identity_and_conflicting_races(
    sql_connection, integration_settings, monkeypatch,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'warehouse_code': 'WH-RACE', 'name': 'Race Warehouse',
            'is_default': False, 'negative_stock_policy': 'WARN',
            'status': 'ACTIVE',
        }
        original_plan = warehouse_provisioning.plan_warehouse

        def synchronized_plan():
            arrived = 0
            gate = asyncio.Event()

            async def value(*args, **kwargs):
                nonlocal arrived
                plan = await original_plan(*args, **kwargs)
                arrived += 1
                if arrived == 2:
                    gate.set()
                await asyncio.wait_for(gate.wait(), timeout=5)
                return plan

            return value

        try:
            async with database.session_factory() as db:
                with pytest.raises(
                    warehouse_provisioning.WarehouseScopeNotFoundError
                ):
                    await original_plan(
                        db, **{**arguments,
                               'organization_id': scope.organization_id + 999999},
                    )
                for changes in (
                    {'status': 'BROKEN'},
                    {'negative_stock_policy': 'IGNORE'},
                    {'warehouse_code': 'invalid code'},
                ):
                    with pytest.raises(
                        warehouse_provisioning.WarehouseProvisioningError
                    ):
                        await original_plan(db, **{**arguments, **changes})

            monkeypatch.setattr(
                warehouse_provisioning, 'plan_warehouse', synchronized_plan(),
            )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    warehouse_provisioning.provision_warehouse(
                        first_db, **arguments,
                    ),
                    warehouse_provisioning.provision_warehouse(
                        second_db, **arguments,
                    ),
                )

            monkeypatch.setattr(
                warehouse_provisioning, 'plan_warehouse', synchronized_plan(),
            )
            conflict_arguments = {
                **arguments, 'warehouse_code': 'WH-CONFLICT',
            }
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                conflicts = await asyncio.gather(
                    warehouse_provisioning.provision_warehouse(
                        first_db, **conflict_arguments,
                    ),
                    warehouse_provisioning.provision_warehouse(
                        second_db, **{
                            **conflict_arguments, 'name': 'Different Warehouse',
                        },
                    ),
                    return_exceptions=True,
                )
            assert sum(
                isinstance(value, warehouse_provisioning.WarehouseConflictError)
                for value in conflicts
            ) == 1

            monkeypatch.setattr(
                warehouse_provisioning, 'plan_warehouse', original_plan,
            )
            async with database.session_factory() as db:
                with pytest.raises(
                    warehouse_provisioning.WarehouseConflictError
                ):
                    await original_plan(
                        db, **{**arguments, 'is_default': True},
                    )
            return first.warehouse.id, second.warehouse.id, {
                first.operation, second.operation,
            }
        finally:
            monkeypatch.setattr(
                warehouse_provisioning, 'plan_warehouse', original_plan,
            )
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'warehouses', scope.tenant_id) == 2
    for table in (
        'stock_movements', 'inventory_lots', 'inventory_cost_layers',
        'inventory_transfers', 'goods_receipts', 'inventory_losses',
        'physical_counts', 'preparation_batches',
        'inventory_valuation_snapshots',
    ):
        assert _count(connection, table, scope.tenant_id) == 0, table


def test_consumption_import_replaces_complete_aggregate_without_inventory_evidence(
    client, sql_connection,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='CONSUMPTION', consumption=True,
    )
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert not result['errors'], result['errors']
    for group in ('consumption_definitions', 'consumption_components'):
        assert result['groups'][group]['classification'] == 'IMPORTABLE_NOW'
        assert result['groups'][group]['authority'] == (
            'restaurant.inventory.consumption_provisioning.'
            'provision_consumption_definition'
        )
        assert result['groups'][group]['created'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT d.id,d.product_id,d.location_id,d.version,d.status,'
            'd.tracking_mode,c.quantity,i.code AS item_code '
            'FROM product_consumption_definitions d '
            'JOIN product_consumption_components c ON c.definition_id=d.id '
            'JOIN inventory_items i ON i.id=c.inventory_item_id '
            'WHERE d.tenant_id=%s', (scope.tenant_id,),
        )
        aggregate = cursor.fetchone()
        cursor.execute(
            'SELECT vc.source_quantity,vc.source_uom,vc.conversion_revision_id,'
            'vc.conversion_factor FROM product_consumption_version_components vc '
            'JOIN product_consumption_versions v ON v.id=vc.version_id '
            'WHERE v.definition_id=%s', (aggregate['id'],),
        )
        evidence = cursor.fetchone()
    assert aggregate == {
        'id': aggregate['id'], 'product_id': aggregate['product_id'],
        'location_id': scope.location_id, 'version': 1, 'status': 'ACTIVE',
        'tracking_mode': 'DERIVABLE', 'quantity': Decimal('2000.000000'),
        'item_code': 'CONSUMPTION',
    }
    assert evidence == {
        'source_quantity': Decimal('2.000000'), 'source_uom': 'BAG',
        'conversion_revision_id': evidence['conversion_revision_id'],
        'conversion_factor': Decimal('1000.000000000000'),
    }
    assert evidence['conversion_revision_id'] is not None

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'product_consumption_definitions', scope.tenant_id) == 1
    assert _count(connection, 'product_consumption_components', scope.tenant_id) == 1
    assert _count(connection, 'product_consumption_versions', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['24_Consumption_Lines']['D2'] = '3.000000'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(client, headers, changed).json()['dataset_fingerprint']
    updated = _confirm(client, headers, changed, scope.location_id, changed_fingerprint)
    assert updated.status_code == 201, updated.text
    for group in ('consumption_definitions', 'consumption_components'):
        assert updated.json()['groups'][group]['updated'] == 1
    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT version FROM product_consumption_definitions WHERE id=%s',
            (aggregate['id'],),
        )
        assert cursor.fetchone()['version'] == 2
        cursor.execute(
            'SELECT quantity FROM product_consumption_components '
            'WHERE definition_id=%s', (aggregate['id'],),
        )
        assert cursor.fetchone()['quantity'] == Decimal('3000.000000')
    assert _count(connection, 'product_consumption_components', scope.tenant_id) == 1
    assert _count(connection, 'product_consumption_versions', scope.tenant_id) == 2
    assert _count(
        connection, 'product_consumption_version_components', scope.tenant_id,
    ) == 2
    for table in (
        'stock_movements', 'inventory_lots', 'inventory_cost_layers',
        'inventory_transfers', 'goods_receipts', 'inventory_losses',
        'physical_counts', 'preparation_batches',
        'inventory_valuation_snapshots',
    ):
        assert _count(connection, table, scope.tenant_id) == 0, table


def test_consumption_authority_validates_evidence_and_converges_concurrently(
    client, sql_connection, integration_settings, monkeypatch,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='CONS-RACE', product=True,
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '07_Products', (
        'PROD-CONS-CONFLICT', 'ORG', None,
        'Product CONS-CONFLICT', None, 'ACTIVE',
    ), row=3)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    seeded = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert seeded.status_code == 201, seeded.text

    async def exercise() -> tuple[int, int, set[str]]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = product_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        component = consumption_provisioning.ConsumptionComponentDefinition(
            inventory_item_code='CONS-RACE', quantity=Decimal('1.000000'),
            uom='BAG',
        )
        arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': scope.location_id,
            'binding_namespace': namespace,
            'product_key': 'PROD-CONS-RACE',
            'tracking_mode': 'DERIVABLE',
            'effective_from': datetime(2030, 1, 1),
            'status': 'ACTIVE', 'components': (component,),
            'actor_id': scope.membership_id,
        }
        original_plan = consumption_provisioning.plan_consumption_definition

        def synchronized_plan():
            arrived = 0
            gate = asyncio.Event()

            async def value(*args, **kwargs):
                nonlocal arrived
                plan = await original_plan(*args, **kwargs)
                arrived += 1
                if arrived == 2:
                    gate.set()
                await asyncio.wait_for(gate.wait(), timeout=5)
                return plan

            return value

        try:
            async with database.session_factory() as db:
                with pytest.raises(
                    consumption_provisioning.ConsumptionScopeNotFoundError
                ):
                    await original_plan(
                        db, **{
                            key: value for key, value in arguments.items()
                            if key not in {'actor_id', 'organization_id'}
                        }, organization_id=scope.organization_id + 999999,
                    )
                invalid_components = (
                    consumption_provisioning.ConsumptionComponentDefinition(
                        'CONS-RACE', Decimal('1.0000001'), 'BAG',
                    ),
                    consumption_provisioning.ConsumptionComponentDefinition(
                        'CONS-RACE', Decimal('1.000000'), 'CASE',
                    ),
                    consumption_provisioning.ConsumptionComponentDefinition(
                        'UNKNOWN', Decimal('1.000000'), 'G',
                    ),
                )
                expected_errors = (
                    consumption_provisioning.ConsumptionProvisioningError,
                    consumption_provisioning.ConsumptionProvisioningError,
                    consumption_provisioning.ConsumptionScopeNotFoundError,
                )
                for invalid, expected_error in zip(
                    invalid_components, expected_errors, strict=True,
                ):
                    with pytest.raises(expected_error):
                        await original_plan(
                            db, **{
                                key: value for key, value in arguments.items()
                                if key not in {'actor_id', 'components'}
                            }, components=(invalid,),
                        )

            monkeypatch.setattr(
                consumption_provisioning, 'plan_consumption_definition',
                synchronized_plan(),
            )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first, second = await asyncio.gather(
                    consumption_provisioning.provision_consumption_definition(
                        first_db, **arguments,
                    ),
                    consumption_provisioning.provision_consumption_definition(
                        second_db, **arguments,
                    ),
                )

            monkeypatch.setattr(
                consumption_provisioning, 'plan_consumption_definition',
                synchronized_plan(),
            )
            conflicting = {
                **arguments, 'product_key': 'PROD-CONS-CONFLICT',
            }
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                conflicts = await asyncio.gather(
                    consumption_provisioning.provision_consumption_definition(
                        first_db, **conflicting,
                    ),
                    consumption_provisioning.provision_consumption_definition(
                        second_db, **{
                            **conflicting,
                            'components': (
                                consumption_provisioning.
                                ConsumptionComponentDefinition(
                                    'CONS-RACE', Decimal('2.000000'), 'BAG',
                                ),
                            ),
                        },
                    ),
                    return_exceptions=True,
                )
            assert sum(
                isinstance(
                    value, consumption_provisioning.ConsumptionConflictError,
                )
                for value in conflicts
            ) == 1
            return first.definition.id, second.definition.id, {
                first.operation, second.operation,
            }
        finally:
            monkeypatch.setattr(
                consumption_provisioning, 'plan_consumption_definition',
                original_plan,
            )
            await database.dispose()

    first_id, second_id, operations = asyncio.run(exercise())
    assert first_id == second_id
    assert operations == {'CREATE', 'UNCHANGED'}
    assert _count(connection, 'product_consumption_definitions', scope.tenant_id) == 2
    assert _count(connection, 'product_consumption_components', scope.tenant_id) == 2
    assert _count(connection, 'product_consumption_versions', scope.tenant_id) == 2
    for table in (
        'stock_movements', 'inventory_lots', 'inventory_cost_layers',
        'inventory_transfers', 'goods_receipts', 'inventory_losses',
        'physical_counts', 'preparation_batches',
        'inventory_valuation_snapshots',
    ):
        assert _count(connection, table, scope.tenant_id) == 0, table


def test_fiscal_import_replays_updates_and_preserves_runtime_resolution(
    client, sql_connection, integration_settings,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='FISCAL', inventory=False,
        fiscal=True,
    )
    analyzed = _analyze(client, headers, content)
    assert analyzed.status_code == 200, analyzed.text
    assert analyzed.json()['analysis']['status'] == 'VALID'
    fingerprint = analyzed.json()['dataset_fingerprint']

    first = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert first.status_code == 201, first.text
    result = first.json()
    assert not result['errors'], result['errors']
    for group, authority in (
        ('tax_rules', 'restaurant.tax.provisioning.provision_tax_rule'),
        (
            'product_fiscal_classifications',
            'restaurant.fiscal_product.provisioning.'
            'provision_product_fiscal_classification',
        ),
    ):
        assert result['groups'][group]['classification'] == 'IMPORTABLE_NOW'
        assert result['groups'][group]['authority'] == authority
        assert result['groups'][group]['created'] == 1
    assert result['required_deferred_groups'] == []

    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT p.id,p.tax_classification_code FROM products p '
            'JOIN product_external_mappings m ON m.product_id=p.id '
            'WHERE p.tenant_id=%s AND m.external_product_id=%s',
            (scope.tenant_id, 'PROD-FISCAL'),
        )
        product = cursor.fetchone()
        cursor.execute(
            'SELECT id,location_id,tax_rate,effective_from,effective_to,status '
            'FROM restaurant_tax_rules WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        rule = cursor.fetchone()
        cursor.execute(
            'SELECT id,product_id,fiscal_jurisdiction_code,'
            'product_classification_code,unit_classification_code,status '
            'FROM product_fiscal_classifications WHERE tenant_id=%s',
            (scope.tenant_id,),
        )
        classification = cursor.fetchone()
    assert product['tax_classification_code'] == 'CLASS-FISCAL'
    assert rule['location_id'] is None
    assert rule['tax_rate'] == Decimal('0.160000')
    assert rule['effective_to'] is None and rule['status'] == 'ACTIVE'
    assert classification == {
        'id': classification['id'], 'product_id': product['id'],
        'fiscal_jurisdiction_code': 'MX',
        'product_classification_code': '50192701',
        'unit_classification_code': 'H87', 'status': 'ACTIVE',
    }

    async def resolve_runtime() -> tuple[Decimal, str]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        try:
            async with database.session_factory() as db:
                tax = await resolve_tax_evidence(db, RestaurantTaxLineCandidate(
                    tenant_id=scope.tenant_id,
                    organization_id=scope.organization_id,
                    location_id=scope.location_id, product_id=product['id'],
                    product_tax_classification_code='CLASS-FISCAL',
                    effective_at=datetime(2026, 6, 1), tax_mode='INCLUDED',
                    quantity=Decimal('1.0000'), unit_price=Decimal('116.0000'),
                    base_amount=Decimal('116.0000'),
                    discount_amount=Decimal('0.0000'),
                    commercial_amount=Decimal('116.0000'),
                ))
                fiscal = await resolve_fiscal_product_evidence(
                    db, FiscalProductClassificationCandidate(
                        tenant_id=scope.tenant_id,
                        organization_id=scope.organization_id,
                        product_id=product['id'], fiscal_jurisdiction_code='MX',
                        effective_at=datetime(2026, 6, 1),
                    ),
                )
                return tax.tax_rate, fiscal.product_classification_code
        finally:
            await database.dispose()

    assert asyncio.run(resolve_runtime()) == (Decimal('0.160000'), '50192701')

    replay = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert replay.status_code == 200 and replay.json()['replay'] is True
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 1
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 1

    workbook = load_workbook(BytesIO(content), data_only=False)
    workbook['28_Tax_Rules']['I2'] = '0.100000'
    workbook['29_Product_Fiscal']['E2'] = '50192702'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    changed = output.getvalue()
    changed_fingerprint = _analyze(
        client, headers, changed,
    ).json()['dataset_fingerprint']
    updated = _confirm(
        client, headers, changed, scope.location_id, changed_fingerprint,
    )
    assert updated.status_code == 201, updated.text
    for group in ('tax_rules', 'product_fiscal_classifications'):
        assert updated.json()['groups'][group]['updated'] == 1
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 1
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 1
    assert _count(connection, 'product_prices', scope.tenant_id) == 0
    assert _count(connection, 'restaurant_order_item_tax_snapshots', scope.tenant_id) == 0
    assert _count(connection, 'restaurant_payments', scope.tenant_id) == 0


def test_fiscal_authorities_reject_invalid_scope_and_converge_concurrently(
    client, sql_connection, integration_settings, monkeypatch,
) -> None:
    connection, prefix = sql_connection
    scope = _prepare(connection, prefix)
    headers = _headers(client, scope)
    content = _workbook(
        f'{prefix}-onboarding', item_code='FISCAL-RACE', inventory=False,
        product=True,
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    _set_row(workbook, '07_Products', (
        'PROD-FISCAL-CONFLICT', 'ORG', None, 'Fiscal Conflict', None, 'ACTIVE',
    ), row=3)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    fingerprint = _analyze(client, headers, content).json()['dataset_fingerprint']
    seeded = _confirm(client, headers, content, scope.location_id, fingerprint)
    assert seeded.status_code == 201, seeded.text

    async def exercise() -> tuple[set[str], int, int, int, int]:
        from app.db.session import DatabaseManager

        database = DatabaseManager(integration_settings)
        namespace = product_provisioning.onboarding_binding_namespace(
            contract_version=CONTRACT_VERSION,
            organization_id=scope.organization_id,
        )
        tax_arguments = {
            'tenant_id': scope.tenant_id,
            'organization_id': scope.organization_id,
            'location_id': None,
            'tax_classification_code': 'CLASS-FISCAL-RACE',
            'jurisdiction_code': 'MX-FEDERAL', 'tax_category': 'IVA',
            'tax_treatment': 'TAXABLE', 'tax_effect': 'TRANSFERRED',
            'tax_rate': Decimal('0.160000'),
            'calculation_policy': 'INCLUDED_PRICE_SINGLE_TAX',
            'rounding_policy': 'DECIMAL_4_HALF_UP',
            'effective_from': datetime(2026, 1, 1), 'effective_to': None,
            'status': 'ACTIVE',
        }

        def synchronize_first_pair(original):
            arrived = 0
            gate = asyncio.Event()

            async def synchronized(*args, **kwargs):
                nonlocal arrived
                result = await original(*args, **kwargs)
                if arrived >= 2:
                    return result
                arrived += 1
                if arrived == 2:
                    gate.set()
                await asyncio.wait_for(gate.wait(), timeout=5)
                return result

            return synchronized

        original_tax_plan = tax_rule_provisioning.plan_tax_rule
        original_fiscal_plan = (
            fiscal_product_provisioning.plan_product_fiscal_classification
        )
        try:
            async with database.session_factory() as db:
                with pytest.raises(tax_rule_provisioning.TaxRuleProvisioningError):
                    await original_tax_plan(
                        db, **{**tax_arguments, 'tax_rate': Decimal('0.1600001')}
                    )
                with pytest.raises(tax_rule_provisioning.TaxRuleProvisioningError):
                    await original_tax_plan(
                        db, **{**tax_arguments, 'status': 'UNKNOWN'}
                    )
                with pytest.raises(tax_rule_provisioning.TaxRuleProvisioningError):
                    await original_tax_plan(
                        db, **{
                            **tax_arguments,
                            'effective_to': datetime(2025, 12, 31),
                        },
                    )
                with pytest.raises(tax_rule_provisioning.TaxRuleScopeError):
                    await original_tax_plan(
                        db, **{
                            **tax_arguments,
                            'organization_id': scope.organization_id + 999999,
                        },
                    )

            monkeypatch.setattr(
                tax_rule_provisioning, 'plan_tax_rule',
                synchronize_first_pair(original_tax_plan),
            )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first_tax, second_tax = await asyncio.gather(
                    tax_rule_provisioning.provision_tax_rule(
                        first_db, **tax_arguments,
                    ),
                    tax_rule_provisioning.provision_tax_rule(
                        second_db, **tax_arguments,
                    ),
                )
            monkeypatch.setattr(
                tax_rule_provisioning, 'plan_tax_rule', original_tax_plan,
            )

            fiscal_arguments = {
                'tenant_id': scope.tenant_id,
                'organization_id': scope.organization_id,
                'location_id': scope.location_id,
                'binding_namespace': namespace,
                'product_key': 'PROD-FISCAL-RACE',
                'tax_classification_code': 'CLASS-FISCAL-RACE',
                'fiscal_jurisdiction_code': 'MX',
                'product_classification_scheme': 'SAT-PRODUCT',
                'product_classification_code': '50192701',
                'unit_classification_scheme': 'SAT-UNIT',
                'unit_classification_code': 'H87',
                'effective_from': datetime(2026, 1, 1),
                'effective_to': None, 'status': 'ACTIVE',
            }
            async with database.session_factory() as db:
                with pytest.raises(
                    fiscal_product_provisioning.FiscalClassificationScopeError
                ):
                    await original_fiscal_plan(
                        db, **{
                            **fiscal_arguments, 'product_key': 'UNKNOWN-PRODUCT',
                        },
                    )
                with pytest.raises(
                    fiscal_product_provisioning.FiscalClassificationScopeError
                ):
                    await original_fiscal_plan(
                        db, **{
                            **fiscal_arguments,
                            'tax_classification_code': 'UNKNOWN-TAX',
                        },
                    )

            monkeypatch.setattr(
                fiscal_product_provisioning,
                'plan_product_fiscal_classification',
                synchronize_first_pair(original_fiscal_plan),
            )
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                first_fiscal, second_fiscal = await asyncio.gather(
                    fiscal_product_provisioning.provision_product_fiscal_classification(
                        first_db, **fiscal_arguments,
                    ),
                    fiscal_product_provisioning.provision_product_fiscal_classification(
                        second_db, **fiscal_arguments,
                    ),
                )
            async with database.session_factory() as db:
                with pytest.raises(tax_rule_provisioning.TaxRuleConflictError):
                    await original_tax_plan(
                        db, **{
                            **tax_arguments,
                            'effective_from': datetime(2026, 6, 1),
                        },
                    )
                with pytest.raises(
                    fiscal_product_provisioning.FiscalClassificationConflictError
                ):
                    await original_fiscal_plan(
                        db, **{
                            **fiscal_arguments,
                            'effective_from': datetime(2026, 6, 1),
                        },
                    )

            monkeypatch.setattr(
                tax_rule_provisioning, 'plan_tax_rule',
                synchronize_first_pair(original_tax_plan),
            )
            conflicting_tax = {
                **tax_arguments,
                'tax_classification_code': 'CLASS-FISCAL-CONFLICT',
            }
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                tax_conflicts = await asyncio.gather(
                    tax_rule_provisioning.provision_tax_rule(
                        first_db, **conflicting_tax,
                    ),
                    tax_rule_provisioning.provision_tax_rule(
                        second_db, **{
                            **conflicting_tax, 'tax_rate': Decimal('0.100000'),
                        },
                    ),
                    return_exceptions=True,
                )
            monkeypatch.setattr(
                tax_rule_provisioning, 'plan_tax_rule', original_tax_plan,
            )

            monkeypatch.setattr(
                fiscal_product_provisioning,
                'plan_product_fiscal_classification',
                synchronize_first_pair(original_fiscal_plan),
            )
            conflicting_fiscal = {
                **fiscal_arguments, 'product_key': 'PROD-FISCAL-CONFLICT',
            }
            async with (
                database.session_factory() as first_db,
                database.session_factory() as second_db,
            ):
                fiscal_conflicts = await asyncio.gather(
                    fiscal_product_provisioning.provision_product_fiscal_classification(
                        first_db, **conflicting_fiscal,
                    ),
                    fiscal_product_provisioning.provision_product_fiscal_classification(
                        second_db, **{
                            **conflicting_fiscal,
                            'product_classification_code': '50192702',
                        },
                    ),
                    return_exceptions=True,
                )
            return (
                {first_tax.operation, second_tax.operation},
                first_fiscal.classification.id,
                second_fiscal.classification.id,
                sum(
                    isinstance(value, tax_rule_provisioning.TaxRuleConflictError)
                    for value in tax_conflicts
                ),
                sum(
                    isinstance(
                        value,
                        fiscal_product_provisioning.
                        FiscalClassificationConflictError,
                    )
                    for value in fiscal_conflicts
                ),
            )
        finally:
            monkeypatch.setattr(
                tax_rule_provisioning, 'plan_tax_rule', original_tax_plan,
            )
            monkeypatch.setattr(
                fiscal_product_provisioning,
                'plan_product_fiscal_classification', original_fiscal_plan,
            )
            await database.dispose()

    tax_operations, first_id, second_id, tax_conflicts, fiscal_conflicts = (
        asyncio.run(exercise())
    )
    assert tax_operations == {'CREATE', 'UNCHANGED'}
    assert first_id == second_id
    assert tax_conflicts == fiscal_conflicts == 1
    assert _count(connection, 'restaurant_tax_rules', scope.tenant_id) == 2
    assert _count(connection, 'product_fiscal_classifications', scope.tenant_id) == 2
    assert _count(connection, 'product_prices', scope.tenant_id) == 0
